from __future__ import annotations

import csv
import json
import os
import socket
import sys
import threading
from dataclasses import dataclass, field
from typing import Any, Generator, List, Optional
from urllib.parse import quote

from services.vault_service import VaultService
from services.server_resolver import ServerResolver
from utils.cli import build_cmd

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
AUX_ROOT = os.path.abspath(os.path.join(BASE_DIR, "..", ".."))
AUTOADA_DIR = os.path.join(AUX_ROOT, "AutoADA")
OUT_ROOT = os.path.join(AUTOADA_DIR, "out")

if AUTOADA_DIR not in sys.path:
    sys.path.insert(0, AUTOADA_DIR)

SERVER_RESOLVER = ServerResolver(os.path.join(AUTOADA_DIR, "config"))
SERVER_RESOLVER._path = os.path.join(AUTOADA_DIR, "config", "servers.json")


@dataclass
class ConsultarResult:
    status: str
    message: str
    files: list[str] = field(default_factory=list)
    extra: dict[str, Any] = field(default_factory=dict)


last_consultar_rtu_result: Optional[ConsultarResult] = None
last_empresas_cache: Optional[list[str]] = None

_STOP_REQUESTED = False
_PROC_LOCK = threading.Lock()
_CURRENT_PROC = None


def reset_stop_flag() -> None:
    global _STOP_REQUESTED
    _STOP_REQUESTED = False


def _should_stop() -> bool:
    return _STOP_REQUESTED


def _set_current_proc(proc) -> None:
    global _CURRENT_PROC
    with _PROC_LOCK:
        _CURRENT_PROC = proc


def request_stop() -> None:
    global _STOP_REQUESTED
    _STOP_REQUESTED = True
    with _PROC_LOCK:
        proc = _CURRENT_PROC
    if proc and proc.poll() is None:
        try:
            proc.terminate()
        except Exception:
            try:
                proc.kill()
            except Exception:
                pass


def _summary_line(message: str, variant: str = "info") -> str:
    return f"SUMMARY::{message}|{variant}\n"


def _result_line(payload: dict[str, Any]) -> str:
    return f"RESULT::{json.dumps(payload, ensure_ascii=False)}\n"


def get_empresas() -> list[str]:
    # Similar a otros detectores: ITCO/TRA en hosts itco/tra, REPS/REPP en rep
    try:
        host = socket.gethostname().lower()
        if host.startswith(("itco1", "tra1", "isa1cct5_p", "desktop-n14qm43", "isa1ccwx_p")):
            return ["ITCO", "TRA"]
        if host.startswith(("rep1", "rep2")):
            return ["REPS", "REPP"]
    except Exception:
        pass
    return ["ITCO", "TRA"]


def get_dominios() -> list[str]:
    return ["CC"]


def _run_subprocess_stream(cmd: list[str], label: str, env: dict[str, str], cwd: str) -> Generator[str, None, int]:
    yield f"\n--- {label} ---\n"
    yield f"$ {' '.join(cmd)}\n"
    import subprocess

    proc = subprocess.Popen(
        cmd,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        bufsize=1,
        cwd=cwd,
        env=env,
    )
    try:
        assert proc.stdout is not None
        for raw in proc.stdout:
            if _should_stop():
                try:
                    proc.terminate()
                except Exception:
                    try:
                        proc.kill()
                    except Exception:
                        pass
                yield f"[{label}] Proceso detenido por el usuario.\n"
                break
            line = raw.rstrip("\r\n")
            if line:
                yield f"[{label}] {line}\n"
    finally:
        try:
            if proc.stdout:
                proc.stdout.close()
        except Exception:
            pass
        _set_current_proc(None)
    rc = proc.wait()
    yield f"[{label}] Código de salida: {rc}\n"
    return rc


def _scada_path(root: str, empresa: str, dominio: str | None = None) -> str:
    suffix = f"{dominio}SCADA" if dominio else "SCADA"
    return os.path.join(root, "out", empresa, suffix)


def _resolve_scada_dir(empresa: str, dominio: str | None = None) -> str:
    """
    Devuelve la carpeta SCADA existente según dominio, con fallback a variantes conocidas.
    Prioriza <dominio>SCADA, luego minúsculas, y finalmente SCADA legacy.
    """
    empresa = (empresa or "").strip().upper()
    candidates: list[str] = []
    if dominio:
        suffix = f"{dominio}SCADA"
        candidates.append(os.path.join(OUT_ROOT, empresa, suffix))
        candidates.append(os.path.join(OUT_ROOT, empresa, suffix.lower()))
    candidates.append(os.path.join(OUT_ROOT, empresa, "SCADA"))

    for candidate in candidates:
        if os.path.isdir(candidate):
            return candidate
    # Si no existe ninguna, devolver la primera esperada (aunque no exista) para mantener comportamiento anterior
    return candidates[0]


def actualizar_rtu_dataset(empresa: str, dominio: str | None = None) -> Generator[str, None, None]:
    """Importa/convierte SCADA (perfil default, solo SCADA) para consultar RTU."""
    empresa = (empresa or "").strip().upper()
    dominio = "CC"
    try:
        env = VaultService.build_env()
    except Exception as exc:
        yield _result_line({"status": "ERROR", "message": f"No se pudo construir el entorno: {exc}"})
        return

    servidor = SERVER_RESOLVER.generar_server(empresa, dominio)
    if not servidor:
        yield _result_line({"status": "ERROR", "message": "No se pudo resolver el servidor para la empresa seleccionada."})
        return

    cmd_import = build_cmd("scripts.importar_all", servidor, empresa, "sca", "--usecase", "default", "--dominio", dominio)
    rc_import = yield from _run_subprocess_stream(cmd_import, "IMPORT-SCADA", env, AUTOADA_DIR)
    if rc_import != 0:
        yield _result_line({"status": "ERROR", "message": f"Importación SCADA falló (rc={rc_import})."})
        return

    cmd_convert = build_cmd("scripts.Convertir_all", empresa, "Validar_HSH", "--only", "sca", "--dominio", dominio)
    rc_convert = yield from _run_subprocess_stream(cmd_convert, "CONVERT-SCADA", env, AUTOADA_DIR)
    if rc_convert != 0:
        yield _result_line({"status": "ERROR", "message": f"Conversión SCADA falló (rc={rc_convert})."})
        return

    yield _result_line({"status": "SUCCESS", "message": "Datos SCADA actualizados."})


def load_rtus(empresa: str, dominio: str | None = None, search: str | None = None, limit: int = 500) -> dict[str, Any]:
    """Lee RTUs/SAS desde SCADA (32_6.csv) y devuelve lista para UI."""
    empresa = (empresa or "").strip().upper()
    scada_dir = _resolve_scada_dir(empresa, dominio)
    file_path = os.path.join(scada_dir, "32_6.csv")
    items: list[str] = []
    if not os.path.isfile(file_path):
        return {"empresa": empresa, "count": 0, "rtus": []}
    def _read_with_encoding(enc: str):
        with open(file_path, "r", encoding=enc, errors="ignore") as fh:
            reader = csv.reader(fh)
            header = next(reader, [])
            header_lower = [h.strip().lower() for h in header]

            def _idx(*names):
                for n in names:
                    if n in header_lower:
                        return header_lower.index(n)
                return None

            rtu_idx = _idx("rtu/sas", "rtu", "#record")
            name_idx = _idx("name", "nombre", "rtu_abbrev")
            for row in reader:
                if rtu_idx is None or rtu_idx >= len(row):
                    continue
                rtu = (row[rtu_idx] or "").strip()
                if not rtu:
                    continue
                name = (row[name_idx] or "").strip() if name_idx is not None and name_idx < len(row) else ""
                label = f"{rtu}: {name}" if name else rtu
                items.append(label)

    try:
        _read_with_encoding("utf-8-sig")
        if not items:
            _read_with_encoding("cp1252")
        if not items:
            # último intento con delimitador ';'
            with open(file_path, "r", encoding="cp1252", errors="ignore") as fh:
                reader = csv.reader(fh, delimiter=";")
                header = next(reader, [])
                header_lower = [h.strip().lower() for h in header]

                def _idx2(*names):
                    for n in names:
                        if n in header_lower:
                            return header_lower.index(n)
                    return None

                rtu_idx = _idx2("rtu/sas", "rtu", "#record")
                name_idx = _idx2("name", "nombre", "rtu_abbrev")
                for row in reader:
                    if rtu_idx is None or rtu_idx >= len(row):
                        continue
                    rtu = (row[rtu_idx] or "").strip()
                    if not rtu:
                        continue
                    name = (row[name_idx] or "").strip() if name_idx is not None and name_idx < len(row) else ""
                    label = f"{rtu}: {name}" if name else rtu
                    items.append(label)
    except Exception:
        return {"empresa": empresa, "count": 0, "rtus": []}

    # de-dup
    seen = set()
    uniq: list[str] = []
    for it in items:
        if it not in seen:
            seen.add(it)
            uniq.append(it)

    if search:
        s = search.strip().lower()
        uniq = [u for u in uniq if s in u.lower()]

    return {"empresa": empresa, "count": len(uniq), "rtus": uniq[:limit], "has_more": len(uniq) > limit}


def consultar_rtu_pipeline(empresa: str, selected_rtus: list[str], dominio: str | None = None) -> Generator[str, None, None]:
    """Ejecuta scripts.consultar_rtu con la lista seleccionada."""
    global last_consultar_rtu_result
    last_consultar_rtu_result = None
    reset_stop_flag()
    empresa = (empresa or "").strip().upper()
    if not empresa:
        yield _result_line({"status": "ERROR", "message": "Debes seleccionar una empresa."})
        return
    if not selected_rtus:
        yield _result_line({"status": "ERROR", "message": "Selecciona al menos una RTU/SAS."})
        return

    try:
        env = VaultService.build_env()
    except Exception as exc:
        yield _result_line({"status": "ERROR", "message": f"No se pudo construir el entorno: {exc}"})
        return

    yield _summary_line(f"{empresa} ({dominio or 'SCADA'}): consulta iniciada", "info")

    rtus_arg = ", ".join(selected_rtus)
    cmd_args = ["--empresa", empresa, "--rtus", rtus_arg]
    if dominio:
        cmd_args.extend(["--dominio", dominio])
    cmd = build_cmd("scripts.consultar_rtu", *cmd_args)
    report_holder: dict[str, str | None] = {"path": None}

    stream = _run_subprocess_stream(cmd, "CONSULTAR-RTU", env, AUTOADA_DIR)
    rc: int | None = None
    try:
        while True:
            chunk = next(stream)
            line = chunk.strip()
            if line.startswith("[CONSULTAR-RTU]") and "RTU_REPORT:" in line:
                try:
                    _, payload = line.split("RTU_REPORT:", 1)
                    report_holder["path"] = payload.strip()
                except Exception:
                    pass
            yield chunk
    except StopIteration as stop:
        rc = stop.value if isinstance(stop.value, int) else 0

    status = "SUCCESS" if rc == 0 else "ERROR"
    message = "Consulta RTU lista" if rc == 0 else "Consulta RTU terminó con errores."
    yield _summary_line(message, "success" if status == "SUCCESS" else "error")
    files: list[str] = []
    if report_holder.get("path"):
        p = report_holder["path"]
        if p and not os.path.isabs(p):
            p = os.path.join(AUTOADA_DIR, p)
        if p and os.path.isfile(p):
            files.append(os.path.normpath(p))
    last_consultar_rtu_result = ConsultarResult(status=status, message=message, files=files, extra={})
    yield _result_line({"status": status, "message": message, "files": files})


def get_last_consultar_rtu_result() -> ConsultarResult | None:
    return last_consultar_rtu_result


def load_consultar_rtu_result_preview(sheet: str | None = None, limit: int = 500) -> dict[str, Any] | None:
    result = last_consultar_rtu_result
    if result is None:
        return None

    report_path = None
    for f in result.files or []:
        if str(f).lower().endswith((".xlsx", ".xlsm", ".xls")):
            report_path = f
            break
    if report_path and not os.path.isabs(report_path):
        report_path = os.path.normpath(os.path.join(AUTOADA_DIR, report_path))
    if report_path and not os.path.isfile(report_path):
        report_path = None

    base_payload: dict[str, Any] = {
        "status": result.status,
        "message": result.message,
        "files": result.files,
        "details": result.extra.get("details", []) if isinstance(result.extra, dict) else [],
        "sheets": [],
        "active_sheet": None,
        "columns": [],
        "rows": [],
        "total": 0,
        "has_more": False,
        "limit": limit,
        "download_url": None,
    }

    if not report_path:
        return base_payload

    from openpyxl import load_workbook as _lb

    wb = _lb(report_path, read_only=True, data_only=True)
    try:
        sheet_names = list(wb.sheetnames)
        if not sheet_names:
            return base_payload
        active_sheet = sheet if sheet in sheet_names else sheet_names[0]
        ws = wb[active_sheet]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            headers_raw = next(rows_iter)
        except StopIteration:
            base_payload["sheets"] = sheet_names
            base_payload["active_sheet"] = active_sheet
            return base_payload

        headers: list[str] = []
        for idx, header in enumerate(headers_raw or (), start=1):
            if isinstance(header, str):
                clean = header.strip()
                headers.append(clean if clean else f"Columna {idx}")
            elif header is None:
                headers.append(f"Columna {idx}")
            else:
                headers.append(str(header))

        preview_rows: list[dict[str, Any]] = []
        row_count = 0
        has_more = False

        for row in rows_iter:
            row_count += 1
            row_dict: dict[str, Any] = {}
            for col_idx, header in enumerate(headers):
                value = row[col_idx] if col_idx < len(row) else None
                row_dict[header] = value
            if row_count <= limit:
                preview_rows.append(row_dict)
            else:
                has_more = True
                break

        download_url = f"/consultar/rtu/result/download?path={quote(report_path)}"

        base_payload.update(
            {
                "sheets": sheet_names,
                "active_sheet": active_sheet,
                "columns": headers,
                "rows": preview_rows,
                "total": row_count,
                "has_more": has_more,
                "limit": limit,
                "download_url": download_url,
                "report_path": report_path,
            }
        )
        return base_payload
    finally:
        wb.close()
