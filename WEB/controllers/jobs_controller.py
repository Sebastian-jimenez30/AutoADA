from __future__ import annotations

import json
import os
import socket
import sys
import tempfile
import threading
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Generator, List, Optional
from urllib.parse import quote

from openpyxl import load_workbook

from services.vault_service import VaultService
from services.server_resolver import ServerResolver
from utils.cli import build_cmd

SUMMARY_VARIANTS = {"info", "success", "warning", "error"}

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
AUX_ROOT = os.path.abspath(os.path.join(BASE_DIR, "..", ".."))
AUTOADA_DIR = os.path.join(AUX_ROOT, "AutoADA")
OUT_ROOT = os.path.join(AUTOADA_DIR, "out")
LOAD_DIR = os.path.join(OUT_ROOT, "Load")
JOBS_UPLOAD_DIR = os.path.join(OUT_ROOT, "jobs_inputs")

if AUTOADA_DIR not in sys.path:
    sys.path.insert(0, AUTOADA_DIR)

SERVER_RESOLVER = ServerResolver(os.path.join(AUTOADA_DIR, "config"))
# Ajustar ruta explícita al servers.json del árbol AutoADA (el resolver original usa get_resource_path relativo al cwd).
SERVER_RESOLVER._path = os.path.join(AUTOADA_DIR, "config", "servers.json")


@dataclass
class JobsCrearResult:
    status: str
    message: str
    files: list[str] = field(default_factory=list)
    extra: dict[str, Any] = field(default_factory=dict)


last_jobs_crear_result: Optional[JobsCrearResult] = None
last_jobs_eliminar_result: Optional[JobsCrearResult] = None
last_jobs_cambiar_result: Optional[JobsCrearResult] = None

_STOP_REQUESTED = False
_PROC_LOCK = threading.Lock()
_CURRENT_PROC = None


def reset_stop_flag() -> None:
    global _STOP_REQUESTED
    _STOP_REQUESTED = False


def _should_stop() -> bool:
    return _STOP_REQUESTED


def _set_current_proc(proc) -> None:
    global _CURRENT_PROC
    with _PROC_LOCK:
        _CURRENT_PROC = proc


def request_stop() -> None:
    """Marca stop y termina el proceso activo si sigue vivo."""
    global _STOP_REQUESTED
    _STOP_REQUESTED = True
    with _PROC_LOCK:
        proc = _CURRENT_PROC
    if proc and proc.poll() is None:
        try:
            proc.terminate()
        except Exception:
            try:
                proc.kill()
            except Exception:
                pass


def _summary_line(message: str, variant: str = "info") -> str | None:
    clean = (message or "").strip()
    if not clean:
        return None
    normalized = variant.lower()
    if normalized not in SUMMARY_VARIANTS:
        normalized = "info"
    return f"SUMMARY::{clean}|{normalized}\n"


def _result_line(payload: dict[str, Any]) -> str:
    return f"RESULT::{json.dumps(payload, ensure_ascii=False)}\n"


def _normalize_path(path: str) -> str:
    p = os.path.normpath(path)
    if not os.path.isabs(p):
        p = os.path.normpath(os.path.join(AUTOADA_DIR, p))
    return p


def _dedupe_paths(paths: list[str]) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for path in paths:
        if not path:
            continue
        normalized = _normalize_path(path)
        key = os.path.normcase(normalized)
        if key in seen:
            continue
        seen.add(key)
        out.append(normalized)
    return out


def _filter_recent_files(paths: list[str], start_ts: float, slack_s: float = 2.0) -> list[str]:
    if not paths:
        return []
    cutoff = start_ts - max(0.0, slack_s)
    recent: list[str] = []
    for path in paths:
        try:
            if os.path.getmtime(path) >= cutoff:
                recent.append(path)
        except Exception:
            continue
    return recent


def _build_jobs_result_preview(
    result: JobsCrearResult,
    sheet: str | None,
    limit: int,
    download_base: str,
) -> dict[str, Any]:
    files = _dedupe_paths(result.files or [])

    base_payload: dict[str, Any] = {
        "status": result.status,
        "message": result.message,
        "files": files,
        "details": result.extra.get("details", []) if isinstance(result.extra, dict) else [],
        "sheets": [],
        "active_sheet": None,
        "columns": [],
        "rows": [],
        "total": 0,
        "has_more": False,
        "limit": limit,
        "download_url": None,
    }

    excel_files = [f for f in files if f.lower().endswith((".xlsx", ".xlsm", ".xls")) and os.path.isfile(f)]
    csv_files = [f for f in files if f.lower().endswith(".csv") and os.path.isfile(f)]

    datasets: dict[str, dict[str, Any]] = {}
    download_map: dict[str, str] = {}
    used_names: set[str] = set()

    def _unique_name(name: str) -> str:
        if name not in used_names:
            used_names.add(name)
            return name
        idx = 2
        while True:
            candidate = f"{name} ({idx})"
            if candidate not in used_names:
                used_names.add(candidate)
                return candidate
            idx += 1

    for excel_path in excel_files:
        try:
            workbook = load_workbook(excel_path, read_only=True, data_only=True)
        except Exception:
            workbook = None
        if not workbook:
            continue
        try:
            base_name = os.path.splitext(os.path.basename(excel_path))[0]
            sheet_key = _unique_name(base_name)
            sheet_names = list(workbook.sheetnames)
            if not sheet_names:
                continue
            sheet_name = sheet_names[0]
            worksheet = workbook[sheet_name]
            rows_iter = worksheet.iter_rows(values_only=True)
            try:
                headers_raw = next(rows_iter)
            except StopIteration:
                headers_raw = []

            headers: list[str] = []
            for idx, header in enumerate(headers_raw or (), start=1):
                if isinstance(header, str):
                    clean = header.strip()
                    headers.append(clean if clean else f"Columna {idx}")
                elif header is None:
                    headers.append(f"Columna {idx}")
                else:
                    headers.append(str(header))

            preview_rows: list[dict[str, Any]] = []
            row_count = 0
            has_more = False

            for row in rows_iter:
                row_count += 1
                row_dict: dict[str, Any] = {}
                for col_idx, header in enumerate(headers):
                    value = row[col_idx] if col_idx < len(row) else None
                    row_dict[header] = value
                if row_count <= limit:
                    preview_rows.append(row_dict)
                else:
                    has_more = True
                    break

            datasets[sheet_key] = {
                "columns": headers,
                "rows": preview_rows,
                "total": row_count,
                "has_more": has_more,
            }
            download_map[sheet_key] = f"{download_base}?path={quote(excel_path)}"
        finally:
            try:
                workbook.close()
            except Exception:
                pass

    for csv_path in csv_files:
        base_name = os.path.splitext(os.path.basename(csv_path))[0]
        sheet_key = _unique_name(base_name)
        rows: list[dict[str, Any]] = []
        total = 0
        has_more = False
        try:
            with open(csv_path, "r", encoding="utf-8", errors="ignore") as fh:
                for line in fh:
                    total += 1
                    if total <= limit:
                        rows.append({"line": line.rstrip("\r\n")})
                    else:
                        has_more = True
                        break
        except Exception:
            continue

        datasets[sheet_key] = {
            "columns": ["line"],
            "rows": rows,
            "total": total,
            "has_more": has_more,
        }
        download_map[sheet_key] = f"{download_base}?path={quote(csv_path)}"

    if not datasets:
        return base_payload

    sheets = list(datasets.keys())
    active_sheet = sheet if sheet in sheets else sheets[0]
    active = datasets.get(active_sheet, {})

    base_payload.update(
        {
            "sheets": sheets,
            "active_sheet": active_sheet,
            "columns": active.get("columns", []),
            "rows": active.get("rows", []),
            "total": active.get("total", 0),
            "has_more": active.get("has_more", False),
            "download_url": download_map.get(active_sheet),
        }
    )
    return base_payload


def _run_subprocess_stream(
    cmd: list[str],
    label: str,
    env: dict[str, str],
    cwd: str,
) -> Generator[str, None, int]:
    yield f"\n--- {label} ---\n"
    yield f"$ {' '.join(cmd)}\n"
    import subprocess

    proc = subprocess.Popen(
        cmd,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        bufsize=1,
        cwd=cwd,
        env=env,
    )
    _set_current_proc(proc)

    try:
        assert proc.stdout is not None
        for raw in proc.stdout:
            if _should_stop():
                try:
                    proc.terminate()
                except Exception:
                    try:
                        proc.kill()
                    except Exception:
                        pass
                yield f"[{label}] Proceso detenido por el usuario.\n"
                break
            line = raw.rstrip("\r\n")
            if line:
                yield f"[{label}] {line}\n"
    finally:
        try:
            if proc.stdout:
                proc.stdout.close()
        except Exception:
            pass
        _set_current_proc(None)

    rc = proc.wait()
    yield f"[{label}] Código de salida: {rc}\n"
    return rc


def _detect_empresas_jobs() -> list[str]:
    """Replica seguridad desktop: ITCO hosts -> ITCO; REP hosts -> REPS; default ITCO."""
    try:
        host = socket.gethostname().lower()
        if host.startswith(("itco1", "tra1", "isa1cct5_p", "desktop-n14qm43", "isa1ccwx_p")):
            return ["ITCO"]
        if host.startswith(("rep1", "rep2")):
            return ["REPS"]
    except Exception:
        pass
    return ["ITCO"]


def get_empresas() -> list[str]:
    return _detect_empresas_jobs()


def get_dominios() -> list[str]:
    # Dominio fijo para jobs: QA (SCADA)
    return ["QA"]


def crear_senales_pipeline(
    empresa: str,
    actualizar: bool,
    archivo_path: str,
    archivo_nombre: str | None,
    dominio: str | None = None,
) -> Generator[str, None, None]:
    """Flujo web para Jobs -> Crear señales (equivalente al handler desktop)."""
    global last_jobs_crear_result
    last_jobs_crear_result = None
    reset_stop_flag()
    run_started = time.time()
    extra_messages: list[str] = []

    def _store(status: str, message: str, files: list[str] | None = None, extra: dict[str, Any] | None = None):
        global last_jobs_crear_result
        last_jobs_crear_result = JobsCrearResult(
            status=status,
            message=message,
            files=files or [],
            extra=extra or {},
        )

    empresa = (empresa or "").strip().upper()
    archivo_nombre = archivo_nombre or os.path.basename(archivo_path)
    dominio = (dominio or "QA").strip().upper()

    yield f"Iniciando creación de señales para empresa={empresa} dominio={dominio} archivo={archivo_nombre}\n"
    summary_line = _summary_line(f"{empresa}: proceso iniciado")
    if summary_line:
        extra_messages.append(f"{empresa}: proceso iniciado")
        yield summary_line

    if not empresa:
        payload = {"status": "ERROR", "message": "Debes seleccionar una empresa válida."}
        _store(payload["status"], payload["message"])
        yield _result_line(payload)
        return
    if not archivo_path or not os.path.isfile(archivo_path):
        payload = {"status": "ERROR", "message": "No se pudo acceder al archivo de entrada."}
        _store(payload["status"], payload["message"])
        yield _result_line(payload)
        return

    try:
        env = VaultService.build_env()
    except Exception as exc:
        message = f"No se pudo construir el entorno: {exc}"
        _store("ERROR", message)
        yield _result_line({"status": "ERROR", "message": message})
        return

    servidor = SERVER_RESOLVER.generar_server(empresa, dominio)
    if not servidor:
        message = "No se pudo resolver el servidor para la empresa seleccionada."
        _store("ERROR", message)
        yield _result_line({"status": "ERROR", "message": message})
        return

    def _stream_step(label: str, cmd: list[str]) -> int:
        stream = _run_subprocess_stream(cmd, label, env=env, cwd=AUTOADA_DIR)
        rc: int | None = None
        try:
            while True:
                chunk = next(stream)
                yield chunk
        except StopIteration as stop:
            rc = stop.value if isinstance(stop.value, int) else 0
        return rc if rc is not None else 0

    if actualizar:
        cmd_import = build_cmd("scripts.importar_all", servidor, empresa, "sca", "--usecase", "jobs_crear_senales", "--dominio", dominio)
        cmd_convert = build_cmd("scripts.Convertir_all", empresa, "jobs", "--dominio", dominio)

        rc_import = yield from _stream_step("IMPORT-SCADA", cmd_import)
        if rc_import != 0:
            msg = f"Importación SCADA falló (rc={rc_import})."
            extra_messages.append(msg)
            line = _summary_line(msg, "error")
            if line:
                yield line
            payload = {"status": "ERROR", "message": msg}
            _store(payload["status"], payload["message"], extra={"details": extra_messages})
            yield _result_line(payload)
            return

        line = _summary_line("Importación SCADA completada", "success")
        if line:
            extra_messages.append("Importación SCADA completada")
            yield line

        rc_convert = yield from _stream_step("CONVERT-JOBS", cmd_convert)
        if rc_convert != 0:
            msg = f"Conversión jobs falló (rc={rc_convert})."
            extra_messages.append(msg)
            line = _summary_line(msg, "error")
            if line:
                yield line
            payload = {"status": "ERROR", "message": msg}
            _store(payload["status"], payload["message"], extra={"details": extra_messages})
            yield _result_line(payload)
            return

        line = _summary_line("Conversión jobs completada", "success")
        if line:
            extra_messages.append("Conversión jobs completada")
            yield line

    # Validación del Excel de entrada
    cmd_scan = build_cmd("scripts.scan_data", archivo_path, empresa, "--dominio", dominio)
    rc_scan = yield from _stream_step("SCAN-DATA", cmd_scan)
    if rc_scan == 2:
        msg = "Validación de datos fallida."
        extra_messages.append(msg)
        line = _summary_line(msg, "error")
        if line:
            yield line
        # Intentar leer detalles de validacion_errores.txt
        err_file = os.path.join(AUTOADA_DIR, "out", "validacion_errores.txt")
        errores: list[str] = []
        try:
            if os.path.exists(err_file):
                with open(err_file, "r", encoding="utf-8") as f:
                    errores = [ln.strip() for ln in f.readlines() if ln.strip()]
        except Exception:
            pass
        payload = {"status": "ERROR", "message": msg, "details": errores}
        _store(payload["status"], payload["message"], files=[err_file] if errores else [], extra={"details": errores})
        yield _result_line(payload)
        return
    elif rc_scan != 0:
        msg = f"Validación/scan finalizó con errores (rc={rc_scan})."
        extra_messages.append(msg)
        line = _summary_line(msg, "error")
        if line:
            yield line
        payload = {"status": "ERROR", "message": msg}
        _store(payload["status"], payload["message"], extra={"details": extra_messages})
        yield _result_line(payload)
        return
    else:
        line = _summary_line("Validación de datos OK", "success")
        if line:
            extra_messages.append("Validación de datos OK")
            yield line

    # Generación de cargas SCADA
    cmd_scada = build_cmd("scripts.SCADA_S-A", empresa, "--dominio", dominio)
    rc_scada = yield from _stream_step("SCADA-S-A", cmd_scada)

    status = "SUCCESS" if rc_scada == 0 else "ERROR"
    message = "Creación de señales completada." if rc_scada == 0 else f"Creación de señales finalizó con errores (rc={rc_scada})."
    line = _summary_line(message, "success" if status == "SUCCESS" else "error")
    if line:
        extra_messages.append(message)
        yield line

    # Recopilar archivos generados
    files: list[str] = []
    expected = [
        os.path.join(LOAD_DIR, "10_SCADA.csv"),
        os.path.join(LOAD_DIR, "32_FEP.csv"),
        os.path.join(LOAD_DIR, "Senales_with_keys.xlsx"),
    ]
    for path in expected:
        if os.path.isfile(path):
            files.append(os.path.normpath(path))
    # incluir cualquier otro archivo en Load
    if os.path.isdir(LOAD_DIR):
        for fname in os.listdir(LOAD_DIR):
            fpath = os.path.join(LOAD_DIR, fname)
            if os.path.isfile(fpath):
                files.append(os.path.normpath(fpath))

    files = _dedupe_paths(files)
    recent_files = _filter_recent_files(files, run_started)
    files = recent_files or files

    extra_payload: dict[str, Any] = {"details": extra_messages}
    report_path = os.path.join(LOAD_DIR, "Senales_with_keys.xlsx")
    if os.path.isfile(report_path):
        extra_payload["report_path"] = os.path.normpath(report_path)

    _store(status, message, files=files, extra=extra_payload)
    payload = {"status": status, "message": message, "files": files, "details": extra_messages}
    yield _result_line(payload)


def get_last_jobs_crear_result() -> JobsCrearResult | None:
    return last_jobs_crear_result


def load_jobs_crear_result_preview(sheet: str | None = None, limit: int = 500) -> dict[str, Any] | None:
    result = last_jobs_crear_result
    if result is None:
        return None
    return _build_jobs_result_preview(result, sheet, limit, "/jobs/crear/result/download")


def eliminar_senales_pipeline(
    empresa: str,
    actualizar: bool,
    archivo_path: str,
    archivo_nombre: str | None,
    dominio: str | None = None,
) -> Generator[str, None, None]:
    """Flujo web para Jobs -> Eliminar señales (equivale al handler desktop)."""
    global last_jobs_eliminar_result
    last_jobs_eliminar_result = None
    run_started = time.time()
    extra_messages: list[str] = []

    def _store(status: str, message: str, files: list[str] | None = None, extra: dict[str, Any] | None = None):
        global last_jobs_eliminar_result
        last_jobs_eliminar_result = JobsCrearResult(
            status=status,
            message=message,
            files=files or [],
            extra=extra or {},
        )

    empresa = (empresa or "").strip().upper()
    archivo_nombre = archivo_nombre or os.path.basename(archivo_path)
    dominio = (dominio or "QA").strip().upper()

    yield f"Iniciando eliminación de señales para empresa={empresa} dominio={dominio} archivo={archivo_nombre}\n"
    summary_line = _summary_line(f"{empresa}: proceso iniciado")
    if summary_line:
        extra_messages.append(f"{empresa}: proceso iniciado")
        yield summary_line

    if not empresa:
        payload = {"status": "ERROR", "message": "Debes seleccionar una empresa válida."}
        _store(payload["status"], payload["message"])
        yield _result_line(payload)
        return
    if not archivo_path or not os.path.isfile(archivo_path):
        payload = {"status": "ERROR", "message": "No se pudo acceder al archivo de entrada."}
        _store(payload["status"], payload["message"])
        yield _result_line(payload)
        return

    try:
        env = VaultService.build_env()
    except Exception as exc:
        message = f"No se pudo construir el entorno: {exc}"
        _store("ERROR", message)
        yield _result_line({"status": "ERROR", "message": message})
        return

    servidor = SERVER_RESOLVER.generar_server(empresa, dominio)
    if not servidor:
        message = "No se pudo resolver el servidor para la empresa seleccionada."
        _store("ERROR", message)
        yield _result_line({"status": "ERROR", "message": message})
        return

    def _stream_step(label: str, cmd: list[str]) -> int:
        stream = _run_subprocess_stream(cmd, label, env=env, cwd=AUTOADA_DIR)
        rc: int | None = None
        try:
            while True:
                chunk = next(stream)
                yield chunk
        except StopIteration as stop:
            rc = stop.value if isinstance(stop.value, int) else 0
        return rc if rc is not None else 0

    if actualizar:
        cmd_import = build_cmd("scripts.importar_all", servidor, empresa, "sca", "--usecase", "jobs_eliminar_senales", "--dominio", dominio)
        cmd_convert = build_cmd("scripts.Convertir_all", empresa, "jobs", "--dominio", dominio)

        rc_import = yield from _stream_step("IMPORT-SCADA", cmd_import)
        if rc_import != 0:
            msg = f"Importación SCADA falló (rc={rc_import})."
            extra_messages.append(msg)
            line = _summary_line(msg, "error")
            if line:
                yield line
            payload = {"status": "ERROR", "message": msg}
            _store(payload["status"], payload["message"], extra={"details": extra_messages})
            yield _result_line(payload)
            return

        line = _summary_line("Importación SCADA completada", "success")
        if line:
            extra_messages.append("Importación SCADA completada")
            yield line

        rc_convert = yield from _stream_step("CONVERT-JOBS", cmd_convert)
        if rc_convert != 0:
            msg = f"Conversión jobs falló (rc={rc_convert})."
            extra_messages.append(msg)
            line = _summary_line(msg, "error")
            if line:
                yield line
            payload = {"status": "ERROR", "message": msg}
            _store(payload["status"], payload["message"], extra={"details": extra_messages})
            yield _result_line(payload)
            return

        line = _summary_line("Conversión jobs completada", "success")
        if line:
            extra_messages.append("Conversión jobs completada")
            yield line

    # Ejecución de eliminación
    cmd_eliminar = build_cmd("scripts.eliminar_senales_scada", archivo_path, empresa, "--dominio", dominio)
    rc_del = yield from _stream_step("ELIMINAR-SENALES", cmd_eliminar)

    status = "SUCCESS" if rc_del == 0 else "ERROR"
    message = "Eliminación de señales completada." if rc_del == 0 else f"Eliminación finalizó con errores (rc={rc_del})."
    line = _summary_line(message, "success" if status == "SUCCESS" else "error")
    if line:
        extra_messages.append(message)
        yield line

    # Recopilar archivos generados
    delete_dir = os.path.join(AUTOADA_DIR, "out", "Delete")
    files: list[str] = []
    expected = [
        os.path.join(delete_dir, "Delete_scada.csv"),
        os.path.join(delete_dir, "change_key.csv"),
        os.path.join(delete_dir, "Delete_controls.csv"),
    ]
    for path in expected:
        if os.path.isfile(path):
            files.append(os.path.normpath(path))
    if os.path.isdir(delete_dir):
        for fname in os.listdir(delete_dir):
            fpath = os.path.join(delete_dir, fname)
            if os.path.isfile(fpath):
                files.append(os.path.normpath(fpath))

    files = _dedupe_paths(files)
    recent_files = _filter_recent_files(files, run_started)
    files = recent_files or files

    extra_payload: dict[str, Any] = {"details": extra_messages}
    _store(status, message, files=files, extra=extra_payload)
    payload = {"status": status, "message": message, "files": files, "details": extra_messages}
    yield _result_line(payload)


def get_last_jobs_eliminar_result() -> JobsCrearResult | None:
    return last_jobs_eliminar_result


def load_jobs_eliminar_result_preview(sheet: str | None = None, limit: int = 500) -> dict[str, Any] | None:
    result = last_jobs_eliminar_result
    if result is None:
        return None
    return _build_jobs_result_preview(result, sheet, limit, "/jobs/eliminar/result/download")


def cambiar_nombre_pipeline(
    empresa: str,
    actualizar: bool,
    archivo_path: str,
    archivo_nombre: str | None,
    dominio: str | None = None,
) -> Generator[str, None, None]:
    """Flujo web para Jobs -> Cambiar nombre (equivalente al handler desktop)."""
    global last_jobs_cambiar_result
    last_jobs_cambiar_result = None
    run_started = time.time()
    extra_messages: list[str] = []

    def _store(status: str, message: str, files: list[str] | None = None, extra: dict[str, Any] | None = None):
        global last_jobs_cambiar_result
        last_jobs_cambiar_result = JobsCrearResult(
            status=status,
            message=message,
            files=files or [],
            extra=extra or {},
        )

    empresa = (empresa or "").strip().upper()
    archivo_nombre = archivo_nombre or os.path.basename(archivo_path)
    dominio = (dominio or "QA").strip().upper()

    yield f"Iniciando cambio de nombre para empresa={empresa} dominio={dominio} archivo={archivo_nombre}\n"
    line = _summary_line(f"{empresa}: proceso iniciado")
    if line:
        extra_messages.append(f"{empresa}: proceso iniciado")
        yield line

    if not empresa:
        payload = {"status": "ERROR", "message": "Debes seleccionar una empresa válida."}
        _store(payload["status"], payload["message"])
        yield _result_line(payload)
        return
    if not archivo_path or not os.path.isfile(archivo_path):
        payload = {"status": "ERROR", "message": "No se pudo acceder al archivo de entrada."}
        _store(payload["status"], payload["message"])
        yield _result_line(payload)
        return

    try:
        env = VaultService.build_env()
    except Exception as exc:
        message = f"No se pudo construir el entorno: {exc}"
        _store("ERROR", message)
        yield _result_line({"status": "ERROR", "message": message})
        return

    servidor = SERVER_RESOLVER.generar_server(empresa, dominio)
    if not servidor:
        message = "No se pudo resolver el servidor para la empresa seleccionada."
        _store("ERROR", message)
        yield _result_line({"status": "ERROR", "message": message})
        return

    def _stream_step(label: str, cmd: list[str]) -> int:
        stream = _run_subprocess_stream(cmd, label, env=env, cwd=AUTOADA_DIR)
        rc: int | None = None
        try:
            while True:
                chunk = next(stream)
                yield chunk
        except StopIteration as stop:
            rc = stop.value if isinstance(stop.value, int) else 0
        return rc if rc is not None else 0

    if actualizar:
        cmd_import = build_cmd("scripts.importar_all", servidor, empresa, "sca", "--usecase", "jobs_crear_senales", "--dominio", dominio)
        cmd_convert = build_cmd("scripts.Convertir_all", empresa, "jobs", "--dominio", dominio)

        rc_import = yield from _stream_step("IMPORT-SCADA", cmd_import)
        if rc_import != 0:
            msg = f"Importación SCADA falló (rc={rc_import})."
            extra_messages.append(msg)
            line = _summary_line(msg, "error")
            if line:
                yield line
            payload = {"status": "ERROR", "message": msg}
            _store(payload["status"], payload["message"], extra={"details": extra_messages})
            yield _result_line(payload)
            return

        line = _summary_line("Importación SCADA completada", "success")
        if line:
            extra_messages.append("Importación SCADA completada")
            yield line

        rc_convert = yield from _stream_step("CONVERT-JOBS", cmd_convert)
        if rc_convert != 0:
            msg = f"Conversión jobs falló (rc={rc_convert})."
            extra_messages.append(msg)
            line = _summary_line(msg, "error")
            if line:
                yield line
            payload = {"status": "ERROR", "message": msg}
            _store(payload["status"], payload["message"], extra={"details": extra_messages})
            yield _result_line(payload)
            return

        line = _summary_line("Conversión jobs completada", "success")
        if line:
            extra_messages.append("Conversión jobs completada")
            yield line

    # Ejecución de cambio de nombre
    cmd_change = build_cmd("scripts.cambiar_nombre_senales_scada", archivo_path, empresa, "--dominio", dominio)
    rc_change = yield from _stream_step("CAMBIO-NOMBRE", cmd_change)

    status = "SUCCESS" if rc_change == 0 else "ERROR"
    message = "Cambio de nombre completado." if rc_change == 0 else f"Cambio de nombre finalizó con errores (rc={rc_change})."
    line = _summary_line(message, "success" if status == "SUCCESS" else "error")
    if line:
        extra_messages.append(message)
        yield line

    # Recopilar archivos generados
    name_dir = os.path.join(AUTOADA_DIR, "out", "Name")
    files: list[str] = []
    expected = [os.path.join(name_dir, "change_key.csv")]
    for path in expected:
        if os.path.isfile(path):
            files.append(os.path.normpath(path))
    if os.path.isdir(name_dir):
        for fname in os.listdir(name_dir):
            fpath = os.path.join(name_dir, fname)
            if os.path.isfile(fpath):
                files.append(os.path.normpath(fpath))

    files = _dedupe_paths(files)
    recent_files = _filter_recent_files(files, run_started)
    files = recent_files or files

    extra_payload: dict[str, Any] = {"details": extra_messages}
    _store(status, message, files=files, extra=extra_payload)
    payload = {"status": status, "message": message, "files": files, "details": extra_messages}
    yield _result_line(payload)


def get_last_jobs_cambiar_result() -> JobsCrearResult | None:
    return last_jobs_cambiar_result


def load_jobs_cambiar_result_preview(sheet: str | None = None, limit: int = 500) -> dict[str, Any] | None:
    result = last_jobs_cambiar_result
    if result is None:
        return None
    return _build_jobs_result_preview(result, sheet, limit, "/jobs/cambiar-nombre/result/download")
