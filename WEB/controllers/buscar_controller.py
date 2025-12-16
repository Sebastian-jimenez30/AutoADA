from __future__ import annotations

import json
import os
import subprocess
from typing import Any, Generator, Iterable, Sequence, Tuple

from openpyxl import load_workbook

from services.vault_service import VaultService
from utils.cli import build_cmd
from utils.data_checks import find_mode_data_ready

from services.server_resolver import ServerResolver

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
AUTOADA_DIR = os.path.abspath(os.path.join(BASE_DIR, "..", "..", "AutoADA"))
SERVER_RESOLVER = ServerResolver(os.path.join(AUTOADA_DIR, "config"))
SERVER_RESOLVER._path = os.path.join(AUTOADA_DIR, "config", "servers.json")
OUT_ROOT = os.path.join(AUTOADA_DIR, "out")
RESULT_DIR = os.path.join(OUT_ROOT, "Find_key")
RESULT_FILE = os.path.join(RESULT_DIR, "Find_Key.xlsx")
SUMMARY_VARIANTS = {"info", "success", "warning", "error"}


def get_empresas() -> Iterable[str]:
    return SERVER_RESOLVER.empresa_claves_view().keys()


def get_dominios() -> Iterable[str]:
    return SERVER_RESOLVER.opciones_dominio_view().keys()


def _validar_keys(cadena: str) -> Tuple[list[str], list[str]]:
    keys = [k.strip().upper() for k in (cadena or "").split(",") if k.strip()]
    validas, invalidas = [], []
    for key in keys:
        if "%" in key:
            validas.append(key)
        elif len(key) == 8 and key[:5].isdigit() and key[6:].isdigit():
            validas.append(key)
        else:
            invalidas.append(key)
    return validas, invalidas


def _result_line(status: str, message: str, path: str | None = None) -> str:
    payload = {"status": status, "message": message}
    if path:
        payload["path"] = path
    return f"RESULT::{json.dumps(payload, ensure_ascii=False)}\n"


def _run_subprocess_stream(cmd: list[str], label: str, env: dict[str, str]) -> Generator[str, None, int]:
    yield f"\n--- {label} ---\n"
    yield f"$ {' '.join(cmd)}\n"
    process = subprocess.Popen(
        cmd,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        bufsize=1,
        cwd=AUTOADA_DIR,
        env=env,
    )

    try:
        assert process.stdout is not None
        for raw_line in process.stdout:
            line = raw_line.rstrip("\r\n")
            if not line:
                continue
            stripped = line.strip()
            if stripped.startswith("SUMMARY::"):
                yield stripped + "\n"
                continue
            yield f"[{label}] {line}\n"
    finally:
        try:
            if process.stdout:
                process.stdout.close()
        except Exception:
            pass

    rc = process.wait()
    yield f"[{label}] Código de salida: {rc}\n"
    return rc


def _summary_line(message: str, variant: str = "info") -> str | None:
    text = (message or "").strip()
    if not text:
        return None
    normalized = variant.lower()
    if normalized not in SUMMARY_VARIANTS:
        normalized = "info"
    return f"SUMMARY::{text}|{normalized}\n"


def buscar_key_pipeline(
    empresa: str,
    dominio: str,
    keys: str,
    forzar_actualizacion: bool,
) -> Generator[str, None, None]:
    empresa = (empresa or "").strip().upper()
    dominio = (dominio or "").strip().upper()
    yield f"Iniciando búsqueda de keys para empresa={empresa} dominio={dominio}\n"
    summary_prefix = empresa if empresa else "SIN_EMPRESA"

    def _summary(message: str, variant: str = "info") -> str | None:
        return _summary_line(message, variant)

    line = _summary(f"{summary_prefix}: proceso de búsqueda iniciado")
    if line:
        yield line

    if not empresa:
        yield _result_line("ERROR", "Debes seleccionar una empresa válida.")
        return
    if not dominio:
        yield _result_line("ERROR", "Debes seleccionar un dominio válido.")
        return

    keys_validas, keys_invalidas = _validar_keys(keys)
    if keys_invalidas:
        yield f"Keys inválidas detectadas: {', '.join(keys_invalidas)}\n"
    if not keys_validas:
        yield _result_line("ERROR", "No hay keys válidas para procesar.")
        line = _summary(f"{summary_prefix}: no hay keys válidas para procesar.", "error")
        if line:
            yield line
        return
    line = _summary(f"{summary_prefix}: {len(keys_validas)} key(s) válidas listas para buscar")
    if line:
        yield line

    try:
        env = VaultService.build_env()
    except Exception as exc:  # pragma: no cover
        yield f"Error cargando variables del vault: {exc}\n"
        yield _result_line("ERROR", "No se pudo construir el entorno. Verifica que el vault esté desbloqueado.")
        line = _summary("Error obteniendo credenciales del vault.", "error")
        if line:
            yield line
        return

    ready, details = find_mode_data_ready(AUTOADA_DIR, empresa, dominio)
    needs_update = forzar_actualizacion or (not ready)
    yield f"Datos locales disponibles: {ready} (forzar={forzar_actualizacion})\n"
    yield f"Detalle OUT/SCADA/HSH/ODSTXT: {details}\n"
    if needs_update:
        line = _summary_line(f"{summary_prefix}: se sincronizarán datos locales antes de procesar archivo")
    else:
        line = _summary_line(f"{summary_prefix}: datos locales vigentes, se omite sincronización previa")
    if line:
        yield line
    if needs_update:
        line = _summary(f"{summary_prefix}: se sincronizarán datos locales antes de buscar")
    else:
        line = _summary(f"{summary_prefix}: datos locales vigentes, se omite sincronización previa")
    if line:
        yield line

    usecase = "buscar_keys"

    def _abort(message: str):
        yield message + "\n"
        yield _result_line("ERROR", message)
        line = _summary_line(message, "error")
        if line:
            yield line
        line = _summary(message, "error")
        if line:
            yield line

    if not needs_update:
        cmd = build_cmd("scripts.buscar_key", empresa, ",".join(keys_validas), "--dominio", dominio)
        line = _summary(f"{summary_prefix}: búsqueda en curso ({len(keys_validas)} key(s))")
        if line:
            yield line
        rc = yield from _run_subprocess_stream(cmd, "BUSCAR", env)
        if rc == 0:
            output_file = os.path.join(OUT_ROOT, "Find_key", "Find_Key.xlsx")
            msg = (
                f"Búsqueda completada exitosamente. "
                f"Archivo esperado en: {output_file}"
            )
            matches = _count_result_rows()
            success_summary = (
                f"{summary_prefix}: búsqueda completada ({matches} coincidencias)."
                if matches is not None
                else f"{summary_prefix}: búsqueda completada."
            )
            line = _summary(success_summary, "success")
            if line:
                yield line
            yield _result_line("SUCCESS", msg, output_file)
        else:
            line = _summary(f"{summary_prefix}: búsqueda falló (rc={rc}).", "error")
            if line:
                yield line
            yield _result_line("ERROR", "El comando de búsqueda finalizó con errores.")
        return

    servidor = SERVER_RESOLVER.generar_server(empresa, dominio)
    if not servidor:
        yield from _abort("No se pudo resolver el servidor para la empresa y dominio seleccionados.")
        return

    yield f"Actualizando datos desde el servidor '{servidor}' antes de buscar.\n"

    steps = [
        ("IMPORT-SCA", build_cmd("scripts.importar_all", servidor, empresa, "sca", "--usecase", usecase, "--dominio", dominio)),
        ("CONVERT-SCA", build_cmd("scripts.Convertir_all", empresa, "Buscar_keys", "--only", "sca", "--dominio", dominio)),
        ("IMPORT-HSH", build_cmd("scripts.importar_all", servidor, empresa, "hsh", "--usecase", usecase)),
        ("CONVERT-HSH", build_cmd("scripts.Convertir_all", empresa, "Buscar_keys", "--only", "hsh")),
        ("IMPORT-ODS", build_cmd("scripts.importar_all", servidor, empresa, "ods", "--usecase", usecase)),
        ("CONVERT-ODS", build_cmd("scripts.Convertir_all", empresa, "Buscar_keys", "--only", "ods")),
        ("CONVERT-ODS_CSV", build_cmd("scripts.Convertir_all", empresa, "Buscar_keys", "--only", "ods_csv")),
    ]

    for label, cmd in steps:
        rc = yield from _run_subprocess_stream(cmd, label, env)
        if rc != 0:
            yield from _abort(f"El paso {label} finalizó con errores (rc={rc}).")
            return

    cmd_buscar = build_cmd("scripts.buscar_key", empresa, ",".join(keys_validas), "--dominio", dominio)
    line = _summary(f"{summary_prefix}: búsqueda en curso ({len(keys_validas)} key(s))")
    if line:
        yield line
    rc_buscar = yield from _run_subprocess_stream(cmd_buscar, "BUSCAR", env)
    if rc_buscar == 0:
        output_file = os.path.join(OUT_ROOT, "Find_key", "Find_Key.xlsx")
        msg = (
            "Búsqueda completada exitosamente tras la actualización. "
            f"Archivo generado en: {output_file}"
        )
        matches = _count_result_rows()
        line = _summary(
            f"{summary_prefix}: búsqueda completada tras actualización ({matches} coincidencias)."
            if matches is not None
            else f"{summary_prefix}: búsqueda completada tras actualización.",
            "success",
        )
        if line:
            yield line
        yield _result_line("SUCCESS", msg, output_file)
    else:
        line = _summary(f"{summary_prefix}: búsqueda falló tras actualización (rc={rc_buscar}).", "error")
        if line:
            yield line
        yield _result_line("ERROR", "El comando de búsqueda finalizó con errores.")


def _resolve_result_path() -> str | None:
    """Devuelve la ruta del Excel generado si existe."""
    path = RESULT_FILE
    if os.path.isfile(path):
        return path
    return None


def _count_result_rows() -> int | None:
    """Cuenta filas de resultados (excluyendo encabezado) si el Excel existe."""
    path = _resolve_result_path()
    if not path:
        return None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return None
    try:
        sheet_names = workbook.sheetnames
        if not sheet_names:
            return 0
        worksheet = workbook[sheet_names[0]]
        count = 0
        for row in worksheet.iter_rows(values_only=True, min_row=2):
            # considera cualquier fila con datos en al menos una columna
            if any(cell not in (None, "") for cell in row):
                count += 1
        return count
    finally:
        workbook.close()


def get_result_path() -> str | None:
    """Ruta pública utilizada por los endpoints para facilitar pruebas."""
    return _resolve_result_path()


def load_result_preview(limit: int = 500, sheet: str | None = None) -> dict[str, Any] | None:
    """
    Lee el archivo de resultados y devuelve una vista previa tabular.
    limit controla cuántas filas se devuelven (resto de filas se indica con has_more).
    """
    path = _resolve_result_path()
    if not path:
        return None

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_names: Sequence[str] = tuple(workbook.sheetnames)
        if not sheet_names:
            return {
                "sheets": [],
                "active_sheet": None,
                "columns": [],
                "rows": [],
                "total": 0,
                "has_more": False,
                "limit": limit,
                "path": path,
            }

        active_sheet_name = sheet or sheet_names[0]
        if active_sheet_name not in sheet_names:
            active_sheet_name = sheet_names[0]

        worksheet = workbook[active_sheet_name]
        rows_iter = worksheet.iter_rows(values_only=True)

        try:
            headers_raw = next(rows_iter)
        except StopIteration:
            return {
                "sheets": list(sheet_names),
                "active_sheet": active_sheet_name,
                "columns": [],
                "rows": [],
                "total": 0,
                "has_more": False,
                "limit": limit,
                "path": path,
            }

        headers = []
        for idx, header in enumerate(headers_raw or (), start=1):
            if isinstance(header, str):
                normalized = header.strip()
                headers.append(normalized if normalized else f"Columna {idx}")
            elif header is None:
                headers.append(f"Columna {idx}")
            else:
                headers.append(str(header))

        preview_rows: list[dict[str, Any]] = []
        total = 0
        has_more = False

        for total, row in enumerate(rows_iter, start=1):
            record: dict[str, Any] = {}
            for col_idx, header in enumerate(headers):
                value = row[col_idx] if col_idx < len(row) else None
                record[header] = value

            if total <= limit:
                preview_rows.append(record)
            else:
                has_more = True
                # continue enumerating to know total rows
                # but avoid storing beyond limit
        # Ajustar total real (si no hubo filas, total=0)
        total_rows = total
        if total_rows == 0 and preview_rows:
            total_rows = len(preview_rows)

        return {
            "sheets": list(sheet_names),
            "active_sheet": active_sheet_name,
            "columns": headers,
            "rows": preview_rows,
            "total": total_rows,
            "has_more": has_more,
            "limit": limit,
            "path": path,
        }
    finally:
        workbook.close()


def buscar_keys_pipeline(
    empresa: str,
    dominio: str,
    archivo_path: str,
    archivo_nombre: str | None,
    forzar_actualizacion: bool,
) -> Generator[str, None, None]:
    empresa = (empresa or "").strip().upper()
    dominio = (dominio or "").strip().upper()
    archivo_nombre = archivo_nombre or os.path.basename(archivo_path)

    yield (
        f"Iniciando búsqueda de keys desde archivo para empresa={empresa} dominio={dominio} "
        f"archivo={archivo_nombre}\n"
    )
    summary_prefix = empresa if empresa else "SIN_EMPRESA"
    file_label = archivo_nombre or os.path.basename(archivo_path)

    def _summary(message: str, variant: str = "info") -> str | None:
        return _summary_line(message, variant)

    line = _summary(f"{summary_prefix}: proceso de búsqueda desde archivo iniciado ({file_label})")
    if line:
        yield line

    if not empresa:
        yield _result_line("ERROR", "Debes seleccionar una empresa válida.")
        line = _summary("Proceso detenido: empresa no proporcionada.", "error")
        if line:
            yield line
        return
    if not dominio:
        yield _result_line("ERROR", "Debes seleccionar un dominio válido.")
        line = _summary("Proceso detenido: dominio no proporcionado.", "error")
        if line:
            yield line
        return

    if not archivo_path or not os.path.isfile(archivo_path):
        yield _result_line("ERROR", "Archivo de keys no disponible o inaccesible.")
        line = _summary_line("Archivo de keys no disponible.", "error")
        if line:
            yield line
        return

    _, extension = os.path.splitext(archivo_path)
    extension = (extension or "").lower()
    if extension not in {".xlsx", ".xlsm", ".xls"}:
        yield _result_line("ERROR", "El archivo debe ser un Excel (.xlsx, .xlsm, .xls).")
        line = _summary_line("Formato de archivo inválido. Debe ser Excel.", "error")
        if line:
            yield line
        return

    try:
        env = VaultService.build_env()
    except Exception as exc:  # pragma: no cover
        yield f"Error cargando variables del vault: {exc}\n"
        yield _result_line("ERROR", "No se pudo construir el entorno. Verifica que el vault esté desbloqueado.")
        line = _summary_line("Error obteniendo credenciales del vault.", "error")
        if line:
            yield line
        return

    ready, details = find_mode_data_ready(AUTOADA_DIR, empresa, dominio)
    needs_update = forzar_actualizacion or (not ready)
    yield f"Datos locales disponibles: {ready} (forzar={forzar_actualizacion})\n"
    yield f"Detalle OUT/SCADA/HSH/ODSTXT: {details}\n"

    usecase = "buscar_keys"

    def _abort(message: str):
        yield message + "\n"
        yield _result_line("ERROR", message)

    if not needs_update:
        cmd = build_cmd("scripts.buscar_keys", empresa, archivo_path, "--dominio", dominio)
        line = _summary_line(f"{summary_prefix}: búsqueda desde archivo en curso ({file_label})")
        if line:
            yield line
        rc = yield from _run_subprocess_stream(cmd, "BUSCAR-ARCHIVO", env)
        if rc == 0:
            output_file = os.path.join(OUT_ROOT, "Find_key", "Find_Key.xlsx")
            msg = (
                "Búsqueda completada exitosamente desde archivo. "
                f"Archivo esperado en: {output_file}"
            )
            matches = _count_result_rows()
            summary_msg = (
                f"{summary_prefix}: búsqueda completada desde archivo ({matches} coincidencias)."
                if matches is not None
                else f"{summary_prefix}: búsqueda completada desde archivo."
            )
            line = _summary_line(summary_msg, "success")
            if line:
                yield line
            yield _result_line("SUCCESS", msg, output_file)
        else:
            line = _summary_line(f"{summary_prefix}: búsqueda desde archivo falló (rc={rc}).", "error")
            if line:
                yield line
            yield _result_line("ERROR", "El comando de búsqueda finalizó con errores.")
        return

    servidor = SERVER_RESOLVER.generar_server(empresa, dominio)
    if not servidor:
        yield from _abort("No se pudo resolver el servidor para la empresa y dominio seleccionados.")
        return

    yield f"Actualizando datos desde el servidor '{servidor}' antes de buscar.\n"
    line = _summary(f"{summary_prefix}: iniciando actualización desde {servidor}")
    if line:
        yield line

    steps = [
        ("IMPORT-SCA", build_cmd("scripts.importar_all", servidor, empresa, "sca", "--usecase", usecase, "--dominio", dominio)),
        ("CONVERT-SCA", build_cmd("scripts.Convertir_all", empresa, "Buscar_keys", "--only", "sca", "--dominio", dominio)),
        ("IMPORT-HSH", build_cmd("scripts.importar_all", servidor, empresa, "hsh", "--usecase", usecase)),
        ("CONVERT-HSH", build_cmd("scripts.Convertir_all", empresa, "Buscar_keys", "--only", "hsh")),
        ("IMPORT-ODS", build_cmd("scripts.importar_all", servidor, empresa, "ods", "--usecase", usecase)),
        ("CONVERT-ODS", build_cmd("scripts.Convertir_all", empresa, "Buscar_keys", "--only", "ods")),
        ("CONVERT-ODS_CSV", build_cmd("scripts.Convertir_all", empresa, "Buscar_keys", "--only", "ods_csv")),
    ]

    step_descriptions = {
        "IMPORT-SCA": "Importación SCADA",
        "CONVERT-SCA": "Conversión SCADA",
        "IMPORT-HSH": "Importación HSH",
        "CONVERT-HSH": "Conversión HSH",
        "IMPORT-ODS": "Importación ODS",
        "CONVERT-ODS": "Conversión ODS",
        "CONVERT-ODS_CSV": "Conversión ODS CSV",
    }

    for label, cmd in steps:
        friendly = step_descriptions.get(label, label.replace("-", " ").title())
        line = _summary(f"{summary_prefix}: {friendly} iniciada")
        if line:
            yield line
        rc = yield from _run_subprocess_stream(cmd, label, env)
        if rc != 0:
            line = _summary(f"{summary_prefix}: {friendly} falló (rc={rc}).", "error")
            if line:
                yield line
            yield from _abort(f"El paso {label} finalizó con errores (rc={rc}).")
            return
        line = _summary(f"{summary_prefix}: {friendly} completada", "success")
        if line:
            yield line

    line = _summary(f"{summary_prefix}: datos sincronizados correctamente", "success")
    if line:
        yield line

    cmd_buscar = build_cmd("scripts.buscar_keys", empresa, archivo_path, "--dominio", dominio)
    line = _summary(f"{summary_prefix}: búsqueda desde archivo en curso ({file_label})")
    if line:
        yield line
    rc_buscar = yield from _run_subprocess_stream(cmd_buscar, "BUSCAR-ARCHIVO", env)
    if rc_buscar == 0:
        output_file = os.path.join(OUT_ROOT, "Find_key", "Find_Key.xlsx")
        msg = (
            "Búsqueda completada exitosamente tras la actualización. "
            f"Archivo generado en: {output_file}"
        )
        matches = _count_result_rows()
        line = _summary(
            f"{summary_prefix}: búsqueda completada tras actualización ({matches} coincidencias)."
            if matches is not None
            else f"{summary_prefix}: búsqueda completada tras actualización.",
            "success",
        )
        if line:
            yield line
        yield _result_line("SUCCESS", msg, output_file)
    else:
        line = _summary(f"{summary_prefix}: búsqueda falló tras actualización (rc={rc_buscar}).", "error")
        if line:
            yield line
        yield _result_line("ERROR", "El comando de búsqueda finalizó con errores.")
