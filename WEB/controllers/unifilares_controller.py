from __future__ import annotations

import json
import os
import socket
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Generator, Optional
from urllib.parse import quote

from openpyxl import load_workbook

from services.vault_service import VaultService
from services.server_resolver import ServerResolver
from utils.cli import build_cmd

SUMMARY_VARIANTS = {"info", "success", "warning", "error"}

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
AUX_ROOT = os.path.abspath(os.path.join(BASE_DIR, "..", ".."))
AUTOADA_DIR = os.path.join(AUX_ROOT, "AutoADA")
OUT_ROOT = os.path.join(AUTOADA_DIR, "out")
UNIFILARES_OUT = os.path.join(OUT_ROOT, "Validacion_Unifilares")
UNIFILARES_UPLOAD = os.path.join(OUT_ROOT, "unifilares_inputs")

if AUTOADA_DIR not in sys.path:
    sys.path.insert(0, AUTOADA_DIR)

SERVER_RESOLVER = ServerResolver(os.path.join(AUTOADA_DIR, "config"))
SERVER_RESOLVER._path = os.path.join(AUTOADA_DIR, "config", "servers.json")


@dataclass
class UnifilaresResult:
    status: str
    message: str
    files: list[str] = field(default_factory=list)
    extra: dict[str, Any] = field(default_factory=dict)


last_unifilares_result: Optional[UnifilaresResult] = None


def _summary_line(message: str, variant: str = "info") -> str | None:
    clean = (message or "").strip()
    if not clean:
        return None
    normalized = variant.lower()
    if normalized not in SUMMARY_VARIANTS:
        normalized = "info"
    return f"SUMMARY::{clean}|{normalized}\n"


def _result_line(payload: dict[str, Any]) -> str:
    return f"RESULT::{json.dumps(payload, ensure_ascii=False)}\n"


def _run_subprocess_stream(
    cmd: list[str],
    label: str,
    env: dict[str, str],
    cwd: str,
) -> Generator[str, None, int]:
    yield f"\n--- {label} ---\n"
    yield f"$ {' '.join(cmd)}\n"
    import subprocess

    proc = subprocess.Popen(
        cmd,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        bufsize=1,
        cwd=cwd,
        env=env,
    )

    try:
        assert proc.stdout is not None
        for raw in proc.stdout:
            line = raw.rstrip("\r\n")
            if line:
                yield f"[{label}] {line}\n"
    finally:
        try:
            if proc.stdout:
                proc.stdout.close()
        except Exception:
            pass

    rc = proc.wait()
    yield f"[{label}] Código de salida: {rc}\n"
    return rc


def _detect_empresas_unifilares() -> list[str]:
    """Replica lógica de seguridad: ITCO hosts -> ITCO/TRA; REP -> REPS/REPP; fallback ITCO/TRA."""
    try:
        host = socket.gethostname().lower()
        if host.startswith(("itco1", "tra1", "isa1cct5_p", "desktop-n14qm43", "isa1ccwx_p")):
            return ["ITCO", "TRA"]
        if host.startswith(("rep1", "rep2")):
            return ["REPS", "REPP"]
    except Exception:
        pass
    return ["ITCO", "TRA"]


def get_empresas() -> list[str]:
    return _detect_empresas_unifilares()


def get_dominios() -> list[str]:
    return list(SERVER_RESOLVER.opciones_dominio_view().keys())


def validar_unifilares_pipeline(
    empresa: str,
    dominio: str,
    archivos: list[str],
    actualizar: bool,
) -> Generator[str, None, None]:
    """Flujo web para validar unifilares (equivalente al handler desktop)."""
    global last_unifilares_result
    last_unifilares_result = None
    extra_messages: list[str] = []

    def _store(status: str, message: str, files: list[str] | None = None, extra: dict[str, Any] | None = None):
        global last_unifilares_result
        last_unifilares_result = UnifilaresResult(
            status=status,
            message=message,
            files=files or [],
            extra=extra or {},
        )

    empresa = (empresa or "").strip().upper()
    dominio = (dominio or "").strip().upper()
    archivos = [a for a in archivos if a]

    yield f"Iniciando validación de unifilares para empresa={empresa} dominio={dominio} archivos={len(archivos)}\n"
    line = _summary_line(f"{empresa}: proceso iniciado")
    if line:
        extra_messages.append(f"{empresa}: proceso iniciado")
        yield line

    if not empresa or not dominio:
        payload = {"status": "ERROR", "message": "Debes seleccionar una empresa y dominio válidos."}
        _store(payload["status"], payload["message"])
        yield _result_line(payload)
        return
    if not archivos and not actualizar:
        payload = {"status": "ERROR", "message": "Debes adjuntar al menos un archivo unifilar o ejecutar actualizar BD."}
        _store(payload["status"], payload["message"])
        yield _result_line(payload)
        return

    try:
        env = VaultService.build_env()
    except Exception as exc:
        message = f"No se pudo construir el entorno: {exc}"
        _store("ERROR", message)
        yield _result_line({"status": "ERROR", "message": message})
        return

    servidor = SERVER_RESOLVER.generar_server(empresa, dominio)
    if not servidor:
        message = "No se pudo resolver el servidor para la empresa/dominio seleccionados."
        _store("ERROR", message)
        yield _result_line({"status": "ERROR", "message": message})
        return

    def _stream_step(label: str, cmd: list[str]) -> int:
        stream = _run_subprocess_stream(cmd, label, env=env, cwd=AUTOADA_DIR)
        rc: int | None = None
        try:
            while True:
                chunk = next(stream)
                yield chunk
        except StopIteration as stop:
            rc = stop.value if isinstance(stop.value, int) else 0
        return rc if rc is not None else 0

    if actualizar:
        # Importar SCADA+ODS (perfil unifilares_validar) y convertir ambos
        cmd_import = build_cmd("scripts.importar_all", servidor, empresa, "sca,ods", "--usecase", "unifilares_validar", "--dominio", dominio)
        rc_import = yield from _stream_step("IMPORT-SCADA/ODS", cmd_import)
        if rc_import != 0:
            msg = f"Importación SCADA/ODS falló (rc={rc_import})."
            extra_messages.append(msg)
            line = _summary_line(msg, "error")
            if line:
                yield line
            payload = {"status": "ERROR", "message": msg}
            _store(payload["status"], payload["message"], extra={"details": extra_messages})
            yield _result_line(payload)
            return
        line = _summary_line("Importación SCADA/ODS completada", "success")
        if line:
            extra_messages.append("Importación SCADA/ODS completada")
            yield line

        # Convertir SCADA
        cmd_convert_sca = build_cmd("scripts.Convertir_all", empresa, "Validar_HSH", "--only", "sca", "--dominio", dominio)
        rc_convert_sca = yield from _stream_step("CONVERT-SCADA", cmd_convert_sca)
        if rc_convert_sca != 0:
            msg = f"Conversión SCADA falló (rc={rc_convert_sca})."
            extra_messages.append(msg)
            line = _summary_line(msg, "error")
            if line:
                yield line
            payload = {"status": "ERROR", "message": msg}
            _store(payload["status"], payload["message"], extra={"details": extra_messages})
            yield _result_line(payload)
            return
        line = _summary_line("Conversión SCADA completada", "success")
        if line:
            extra_messages.append("Conversión SCADA completada")
            yield line

        # Convertir ODS
        cmd_convert_ods = build_cmd("scripts.Convertir_all", empresa, "Validar_HSH", "--only", "ods")
        rc_convert_ods = yield from _stream_step("CONVERT-ODS", cmd_convert_ods)
        if rc_convert_ods != 0:
            msg = f"Conversión ODS falló (rc={rc_convert_ods})."
            extra_messages.append(msg)
            line = _summary_line(msg, "error")
            if line:
                yield line
            payload = {"status": "ERROR", "message": msg}
            _store(payload["status"], payload["message"], extra={"details": extra_messages})
            yield _result_line(payload)
            return
        line = _summary_line("Conversión ODS completada", "success")
        if line:
            extra_messages.append("Conversión ODS completada")
            yield line

    # Si solo se pidió actualizar y no hay archivos, terminar aquí
    if actualizar and not archivos:
        status = "SUCCESS"
        message = "Actualización de datos completada."
        line = _summary_line(message, "success")
        if line:
            extra_messages.append(message)
            yield line
        files: list[str] = []
        _store(status, message, files=files, extra={"details": extra_messages})
        payload = {"status": status, "message": message, "files": files, "details": extra_messages}
        yield _result_line(payload)
        return

    # Ejecutar validación de unifilares
    cmd_validar = build_cmd("scripts.Validacion_unifilares", "--archivos", *archivos, "--empresa", empresa, "--dominio", dominio)
    rc_val = yield from _stream_step("VALIDAR-UNIFILARES", cmd_validar)

    status = "SUCCESS" if rc_val == 0 else "ERROR"
    message = "Validación de unifilares completada." if rc_val == 0 else f"Validación finalizó con errores (rc={rc_val})."
    line = _summary_line(message, "success" if status == "SUCCESS" else "error")
    if line:
        extra_messages.append(message)
        yield line

    files: list[str] = []
    if os.path.isdir(UNIFILARES_OUT):
        for fname in os.listdir(UNIFILARES_OUT):
            fpath = os.path.join(UNIFILARES_OUT, fname)
            if os.path.isfile(fpath):
                files.append(os.path.normpath(fpath))

    extra_payload: dict[str, Any] = {"details": extra_messages}
    _store(status, message, files=files, extra=extra_payload)
    payload = {"status": status, "message": message, "files": files, "details": extra_messages}
    yield _result_line(payload)


def get_last_unifilares_result() -> UnifilaresResult | None:
    return last_unifilares_result


def load_unifilares_result_preview(sheet: str | None = None, limit: int = 500) -> dict[str, Any] | None:
    result = last_unifilares_result
    if result is None:
        return None

    def _normalize(path: str) -> str:
        p = os.path.normpath(path)
        if not os.path.isabs(p):
            p = os.path.normpath(os.path.join(AUTOADA_DIR, p))
        return p

    # Preferir Excel de validación
    report_path = None
    for f in result.files or []:
        if str(f).lower().endswith(".xlsx"):
            report_path = f
            break
    if report_path:
        report_path = _normalize(report_path)
    if report_path and not os.path.isfile(report_path):
        report_path = None
    if not report_path:
        return {
            "status": result.status,
            "message": result.message,
            "files": result.files,
            "details": result.extra.get("details", []) if isinstance(result.extra, dict) else [],
            "sheets": [],
            "active_sheet": None,
            "columns": [],
            "rows": [],
            "total": 0,
            "has_more": False,
            "limit": limit,
            "download_url": None,
        }

    wb = load_workbook(report_path, read_only=True, data_only=True)
    try:
        sheet_names = list(wb.sheetnames)
        if not sheet_names:
            return {
                "status": result.status,
                "message": result.message,
                "files": result.files,
                "details": result.extra.get("details", []) if isinstance(result.extra, dict) else [],
                "sheets": [],
                "active_sheet": None,
                "columns": [],
                "rows": [],
                "total": 0,
                "has_more": False,
                "limit": limit,
                "download_url": None,
            }

        active_sheet = sheet if sheet in sheet_names else sheet_names[0]
        ws = wb[active_sheet]
        rows_iter = ws.iter_rows(values_only=True)

        try:
            headers_raw = next(rows_iter)
        except StopIteration:
            return {
                "status": result.status,
                "message": result.message,
                "files": result.files,
                "details": result.extra.get("details", []) if isinstance(result.extra, dict) else [],
                "sheets": sheet_names,
                "active_sheet": active_sheet,
                "columns": [],
                "rows": [],
                "total": 0,
                "has_more": False,
                "limit": limit,
                "download_url": None,
            }

        headers: list[str] = []
        for idx, header in enumerate(headers_raw or (), start=1):
            if isinstance(header, str):
                clean = header.strip()
                headers.append(clean if clean else f"Columna {idx}")
            elif header is None:
                headers.append(f"Columna {idx}")
            else:
                headers.append(str(header))

        preview_rows: list[dict[str, Any]] = []
        row_count = 0
        has_more = False

        for row in rows_iter:
            row_count += 1
            row_dict: dict[str, Any] = {}
            for col_idx, header in enumerate(headers):
                value = row[col_idx] if col_idx < len(row) else None
                row_dict[header] = value
            if row_count <= limit:
                preview_rows.append(row_dict)
            else:
                has_more = True
                break

        download_url = f"/unifilares/validar/result/download?path={quote(report_path)}"

        return {
            "status": result.status,
            "message": result.message,
            "files": result.files,
            "details": result.extra.get("details", []) if isinstance(result.extra, dict) else [],
            "sheets": sheet_names,
            "active_sheet": active_sheet,
            "columns": headers,
            "rows": preview_rows,
            "total": row_count,
            "has_more": has_more,
            "limit": limit,
            "download_url": download_url,
            "report_path": report_path,
        }
    finally:
        wb.close()
