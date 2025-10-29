"""Genera un diccionario de terminaciones de tags a partir de una plantilla Excel."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Dict, List

from openpyxl import load_workbook

try:
    from utils.paths import project_root
except Exception:  # pragma: no cover
    project_root = lambda: Path(__file__).resolve().parents[1]


DEFAULT_SHEET = "Senales_Analogas_Estado_SCADA"
COL_UNIDAD = 6   # columna F
TERMINATION_COLUMNS: Dict[int, str] = {
    7: ".Q",      # columna G
    8: ".State",          # columna H
    9: ".Value",      # columna I
    10: ".Estimated", # columna J
}
START_ROW = 3


def _normalize(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def generar_diccionario(path_excel: Path, hoja: str = DEFAULT_SHEET) -> Dict[str, List[str]]:
    wb = load_workbook(path_excel, data_only=True)
    try:
        ws = wb[hoja]
        diccionario: Dict[str, set[str]] = {}
        for row in range(START_ROW, ws.max_row + 1):
            unidad = _normalize(ws.cell(row=row, column=COL_UNIDAD).value)
            if not unidad:
                continue

            terminaciones = [
                sufijo
                for col, sufijo in TERMINATION_COLUMNS.items()
                if _normalize(ws.cell(row=row, column=col).value).lower() == "x"
            ]
            if not terminaciones:
                continue

            diccionario.setdefault(unidad, set()).update(terminaciones)

        return {k: sorted(v) for k, v in diccionario.items()}
    finally:
        wb.close()


def _default_paths() -> tuple[Path, Path]:
    root = Path(project_root())
    template = root / "templates" / "ProyectoPI_EstandarSenales&TAG.xlsx"
    destino = root / "config" / "diccionario_tags.json"
    return template, destino


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    default_excel, default_out = _default_paths()
    parser.add_argument("--excel", type=Path, default=default_excel,
                        help=f"Ruta del Excel fuente (default: {default_excel})")
    parser.add_argument("--sheet", default=DEFAULT_SHEET,
                        help=f"Nombre de la hoja a leer (default: {DEFAULT_SHEET})")
    parser.add_argument("--output", type=Path, default=default_out,
                        help=f"Ruta del JSON de salida (default: {default_out})")
    return parser.parse_args()


def main() -> int:
    args = _parse_args()
    if not args.excel.exists():
        raise FileNotFoundError(f"No se encontró el archivo Excel: {args.excel}")

    data = generar_diccionario(args.excel, hoja=args.sheet)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8") as fh:
        json.dump(data, fh, indent=4, ensure_ascii=False)

    print(f"Diccionario exportado a {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
