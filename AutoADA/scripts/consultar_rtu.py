import argparse
import logging
import os
import sys
from datetime import datetime
from typing import Any, Dict, List, Tuple, Union

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

try:
    from scripts import _Logger as Logger
except Exception:
    try:
        import _Logger as Logger
    except Exception:
        sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), ".")))
        import _Logger as Logger  # type: ignore

# ---------------------------------------------------------------------------
# Logger global
# ---------------------------------------------------------------------------
RT = os.getcwd()
log_dir = os.path.join(RT, "log")
os.makedirs(log_dir, exist_ok=True)
log_path = os.path.join(log_dir, "consultar_rtu.log")
try:
    logger, logger_console = Logger.initlog(log_path)
except Exception:
    logger = logging.getLogger("consultar_rtu.file")
    logger_console = logging.getLogger("consultar_rtu.console")
    logging.basicConfig(level=logging.INFO)

# ---------------------------------------------------------------------------
# Constantes
# ---------------------------------------------------------------------------
DEFAULT_RTUS = ['106: SABANT', '128: SABA_TRN']
TYPE_MAPPINGS = {
    'status': {1: "T_IND", 2: "T_I&C", 3: "T_CTL"},
    'analog': {1: "T_ANLG", 2: "C_ANLG"},
    'controls': {
        0: 'Auto', 1: 'Raise Only', 2: 'Lower Only', 3: 'Raise-Lower',
        4: 'Open Only', 5: 'Close Only', 6: 'Open-Close', 7: 'Setpoint',
        8: 'Analog', 9: 'Status Out'
    },
    'control_format': {0: 'Default', 1: 'SBO', 2: 'Direct', 3: 'Direct No Ack'},
    'scada_type': {4: 'STATUS', 7: 'SETPOINT'}
}

# ---------------------------------------------------------------------------
# Utilidades internas
# ---------------------------------------------------------------------------
def _parse_rtu_argument(value: Union[str, None], default: List[str]) -> List[str]:
    if value is None:
        return default
    items: List[str] = []
    for raw in value.split(','):
        candidate = raw.strip()
        if candidate:
            items.append(candidate)
    return items or default


def _scada_path(root: str, empresa: str, dominio: str | None = None) -> str:
    """Ruta a SCADA según dominio (default SCADA)."""
    suffix = f"{dominio}SCADA" if dominio else "SCADA"
    return os.path.join(root, "out", empresa, suffix)


def read_dataframes(base_dir: str, empresa: str, dominio: str | None = None) -> Dict[str, pd.DataFrame]:
    Logger.write_log().log_all('info', 'Inicia lectura CSV', logger_console, logger)
    ruta = _scada_path(base_dir, empresa, dominio)
    if not os.path.isdir(ruta):
        raise FileNotFoundError(f"No se encontró el directorio de SCADA: {ruta}")

    file_configs = {
        'df_32_10': {'file': '32_10.csv', 'usecols': ['Key', 'IntParms']},
        'df_32_20': {
            'file': '32_20.csv',
            'usecols': ['pRTU', 'SourceKey', 'point_address', 'proto_parms',
                        'proto_parms[1]', 'control_type', 'control_format',
                        'SourceObject', 'control_bit_parms']
        },
        'df_10_2': {'file': '10_2.csv', 'usecols': ['#record', 'Name']},
        'df_10_4': {'file': '10_4.csv', 'usecols': ['Key', 'Type', 'pStation', 'Name', 'pRTU']},
        'df_10_5': {'file': '10_5.csv', 'usecols': ['Key', 'Type', 'pStation', 'Name', 'pRTU']},
        'df_10_7': {'file': '10_7.csv', 'usecols': ['Key', 'pStation', 'Name'], 'dtype': {'Key': str}}
    }

    dfs: Dict[str, pd.DataFrame] = {}
    for key, cfg in file_configs.items():
        file_path = os.path.join(ruta, cfg['file'])
        Logger.write_log().log_all('info', f"Leyendo archivo {cfg['file']}...", logger_console, logger)
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"No se encontró {cfg['file']} en {ruta}")
        kwargs = {
            'encoding': 'utf-8',
            'sep': ',',
            'usecols': cfg['usecols'],
            'low_memory': False
        }
        if 'dtype' in cfg:
            kwargs['dtype'] = cfg['dtype']
        dfs[key] = pd.read_csv(file_path, **kwargs)
        Logger.write_log().log_all('info', f"Archivo {cfg['file']} cargado. Filas: {len(dfs[key])}", logger_console, logger)

    Logger.write_log().log_all('info', 'CSV cargados correctamente', logger_console, logger)
    return dfs


def crear_diccionarios(dfs: Dict[str, pd.DataFrame]) -> Dict[str, Any]:
    Logger.write_log().log_all('info', 'Creando diccionarios', logger_console, logger)
    diccionarios = {
        'dicc_32_10': dict(zip(
            dfs['df_32_10']['Key'],
            pd.to_numeric(dfs['df_32_10']['IntParms'], errors='coerce').dropna().astype(int)
        )),
        'dicc_10_2': {
            int(rec): f"{int(rec)}:{name}"
            for rec, name in zip(dfs['df_10_2']['#record'], dfs['df_10_2']['Name'])
            if pd.notna(rec) and pd.notna(name)
        },
        'dicc_10_4': {
            str(key).replace('.0', ''): (name, station)
            for key, name, station in zip(dfs['df_10_4']['Key'], dfs['df_10_4']['Name'], dfs['df_10_4']['pStation'])
        },
        'dicc_10_7': {
            str(key).replace('.0', ''): (name, station)
            for key, name, station in zip(dfs['df_10_7']['Key'], dfs['df_10_7']['Name'], dfs['df_10_7']['pStation'])
        }
    }
    for name, dicc in diccionarios.items():
        Logger.write_log().log_all('info', f"Diccionario {name} creado con {len(dicc)} elementos", logger_console, logger)
    return diccionarios


def procesar_rtus(rtu_list: Union[str, List[str]]) -> Tuple[List[int], Dict[int, str]]:
    Logger.write_log().log_all('info', f"Procesando RTUs: {rtu_list}", logger_console, logger)
    rtus = [rtu_list] if isinstance(rtu_list, str) else list(rtu_list)
    rtus_procesadas: List[Dict[str, Union[int, str]]] = []

    for rtu in rtus:
        for sub_rtu in (rtu.split(',') if ',' in rtu else [rtu]):
            if ':' in sub_rtu:
                numero_str, nombre = sub_rtu.split(':', 1)
                numero_str = numero_str.strip()
                if numero_str.isdigit():
                    numero = int(numero_str)
                    rtus_procesadas.append({'numero': numero, 'nombre': nombre.strip()})
                    Logger.write_log().log_all('info', f"RTU procesada: {numero} - {nombre.strip()}", logger_console, logger)

    numeros = [entry['numero'] for entry in rtus_procesadas]
    nombres = {entry['numero']: entry['nombre'] for entry in rtus_procesadas}

    if not rtus_procesadas:
        Logger.write_log().log_all('warning', 'Sin RTU validas para procesar', logger_console, logger)

    return numeros, nombres


def obtener_name_station(row, dicc_10_4, dicc_10_7):
    key = str(row['SCADA Key']).replace('.0', '')
    if row['Scada Type'] == 'STATUS':
        return dicc_10_4.get(key, ('', ''))
    if row['Scada Type'] == 'SETPOINT':
        return dicc_10_7.get(key, ('', ''))
    return '', ''


def procesar_df(df_filtrado, rtu_numero_a_nombre, dicc_32_10, dicc_10_2, type_map):
    if df_filtrado.empty:
        return None
    resultado = pd.DataFrame({
        'RTU/SAS': df_filtrado['pRTU'].map(rtu_numero_a_nombre),
        'IOA': df_filtrado['Key'].map(dicc_32_10),
        'SCADA Key': df_filtrado['Key'],
        'Type': df_filtrado['Type'].map(type_map),
        'Station': df_filtrado['pStation'].map(dicc_10_2),
        'Name': df_filtrado['Name']
    })
    return resultado.sort_values(by=['RTU/SAS', 'IOA'], ascending=[True, True])


def aplicar_formato(sheet, sheet_name):
    sheet.auto_filter.ref = sheet.dimensions
    columnas_colorear = {'STATUS': 2, 'CMD': 3, 'ANALOG': 2}

    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center')

    if sheet_name in columnas_colorear:
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row,
                                  min_col=columnas_colorear[sheet_name],
                                  max_col=columnas_colorear[sheet_name]):
            for cell in row:
                cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                cell.alignment = Alignment(horizontal='center', vertical='center')

    for col in sheet.columns:
        ancho = max((len(str(cell.value)) if cell.value else 0) for cell in col) + 2
        sheet.column_dimensions[col[0].column_letter].width = ancho
        for cell in col:
            cell.alignment = Alignment(horizontal='center', vertical='center')


def procesar_datos_control(df_filtrado, rtu_numero_a_nombre, diccionarios, type_mappings):
    Logger.write_log().log_all('info', 'Procesando datos de control', logger_console, logger)
    if df_filtrado.empty:
        Logger.write_log().log_all('warning', 'Sin datos de control (DataFrame vacio)', logger_console, logger)
        return None

    df_final = pd.DataFrame({
        'RTU/SAS': df_filtrado['pRTU'].map(rtu_numero_a_nombre),
        'SCADA Key': df_filtrado['SourceKey'],
        'ControlAddress': df_filtrado['point_address'].astype(int),
        'Control Parm': df_filtrado['proto_parms'].astype(int),
        'Control Subtype': df_filtrado['proto_parms[1]'].astype(int),
        'Control Type': df_filtrado['control_type'].map(type_mappings['controls']),
        'Control Format': df_filtrado['control_format'].map(type_mappings['control_format']),
        'Scada Type': df_filtrado['SourceObject'].map(type_mappings['scada_type'])
    }).dropna(subset=['SCADA Key'])

    Logger.write_log().log_all('info', 'Aplicando datos de nombre y estacion', logger_console, logger)
    df_final[['Name', 'Station']] = df_final.apply(
        lambda row: pd.Series(obtener_name_station(row, diccionarios['dicc_10_4'], diccionarios['dicc_10_7'])),
        axis=1
    )
    df_final['Station'] = df_final['Station'].map(diccionarios['dicc_10_2'])
    df_final['Reverse Control'] = df_filtrado['control_bit_parms'].astype(int)

    resultado = df_final.sort_values(by=['RTU/SAS', 'ControlAddress'], ascending=[True, True])
    Logger.write_log().log_all('info', f'Procesamiento listo: {len(resultado)} registros de control', logger_console, logger)
    return resultado


def generar_excel(dfs_finales: Dict[str, pd.DataFrame], output_path: str) -> None:
    Logger.write_log().log_all('info', f'Generando archivo Excel en: {output_path}', logger_console, logger)
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        for sheet_name, df in dfs_finales.items():
            if df is not None and not df.empty:
                Logger.write_log().log_all('info', f'Guardando hoja "{sheet_name}" con {len(df)} registros', logger_console, logger)
                df.to_excel(writer, sheet_name=sheet_name, index=False)

    wb = load_workbook(output_path)
    for sheet_name in dfs_finales.keys():
        if sheet_name in wb.sheetnames:
            Logger.write_log().log_all('info', f'Aplicando formato a hoja "{sheet_name}"', logger_console, logger)
            aplicar_formato(wb[sheet_name], sheet_name)
    wb.save(output_path)
    Logger.write_log().log_all('info', f'Archivo Excel guardado: {output_path}', logger_console, logger)


def ejecutar_consulta_rtu(empresa: str, rtu_list: List[str], base_dir: Union[str, None] = None, dominio: str | None = None) -> str:
    if not empresa:
        raise ValueError("Debe proporcionar una empresa válida")
    if not rtu_list:
        raise ValueError("Debe proporcionar al menos una RTU/SAS")

    empresa = empresa.strip().upper()
    base_dir = base_dir or os.getcwd()

    Logger.write_log().log_all('info', '=' * 50, logger_console, logger)
    Logger.write_log().log_all('info', f'Inicio consulta RTU/SAS: {", ".join(rtu_list)}', logger_console, logger)
    Logger.write_log().log_all('info', f'Empresa: {empresa}', logger_console, logger)
    Logger.write_log().log_all('info', '=' * 50, logger_console, logger)

    start_time = datetime.now()
    dfs = read_dataframes(base_dir, empresa, dominio=dominio)
    diccionarios = crear_diccionarios(dfs)
    numeros_rtu, rtu_numero_a_nombre = procesar_rtus(rtu_list)
    if not numeros_rtu:
        raise ValueError("No se pudieron interpretar las RTU/SAS proporcionadas")

    Logger.write_log().log_all('info', 'Paso 4 filtrando datos por RTU', logger_console, logger)
    df_filtrados = {
        key: dfs[key][dfs[key]['pRTU'].isin(numeros_rtu)].copy()
        for key in ['df_10_4', 'df_10_5', 'df_32_20']
    }

    Logger.write_log().log_all('info', 'Paso 5 procesando datos', logger_console, logger)
    dfs_finales = {
        'STATUS': procesar_df(df_filtrados['df_10_4'], rtu_numero_a_nombre,
                              diccionarios['dicc_32_10'], diccionarios['dicc_10_2'], TYPE_MAPPINGS['status']),
        'ANALOG': procesar_df(df_filtrados['df_10_5'], rtu_numero_a_nombre,
                              diccionarios['dicc_32_10'], diccionarios['dicc_10_2'], TYPE_MAPPINGS['analog']),
        'CMD': procesar_datos_control(df_filtrados['df_32_20'], rtu_numero_a_nombre,
                                      diccionarios, TYPE_MAPPINGS)
    }

    for sheet_name, df in dfs_finales.items():
        count = len(df) if df is not None else 0
        Logger.write_log().log_all('info', f'Registros en {sheet_name}: {count}', logger_console, logger)

    Logger.write_log().log_all('info', 'Paso 6 generando Excel', logger_console, logger)
    output_dir = os.path.join(base_dir, 'out', 'Consulta RTU')
    os.makedirs(output_dir, exist_ok=True)
    output_file = os.path.join(output_dir, 'DB_Monarch_Signals.xlsx')
    generar_excel(dfs_finales, output_file)

    end_time = datetime.now()
    Logger.write_log().log_all('info', '=' * 50, logger_console, logger)
    Logger.write_log().log_all('info', 'Proceso finalizado', logger_console, logger)
    Logger.write_log().log_all('info', f'Tiempo de ejecucion: {end_time - start_time}', logger_console, logger)
    Logger.write_log().log_all('info', f'Archivo generado: {output_file}', logger_console, logger)
    Logger.write_log().log_all('info', '=' * 50, logger_console, logger)

    print(f"Archivo generado exitosamente: {output_file}")
    print(f"RTU_REPORT:{output_file}")
    return output_file


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description='Consulta señales SCADA asociadas a RTU/SAS y genera un Excel resumido.')
    parser.add_argument('--empresa', '-e', default=None,
                        help='Nombre de la empresa (por defecto ITCO).')
    parser.add_argument('--rtus', '-r', default=None,
                        help="Lista de RTU/SAS separadas por comas en formato '106: SABANT,128: SABA_TRN'.")
    parser.add_argument('--dominio', '-d', default=None,
                        help='Dominio (ej: CC, QA) para elegir la carpeta SCADA')
    return parser


def main(argv: Union[List[str], None] = None) -> int:
    parser = build_arg_parser()
    args = parser.parse_args(argv)

    empresa = (args.empresa or 'ITCO').strip()
    rtus = _parse_rtu_argument(args.rtus, DEFAULT_RTUS)

    ejecutar_consulta_rtu(empresa, rtus, dominio=args.dominio)
    return 0


if __name__ == '__main__':
    raise SystemExit(main())

