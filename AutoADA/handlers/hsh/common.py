from __future__ import annotations

import base64
import json
import os
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable, Dict, Iterable, List, Optional, Set, Tuple

import tkinter as tk
from tkinter import messagebox

try:
    from openpyxl import load_workbook, Workbook
except Exception:
    load_workbook = None  # type: ignore
    Workbook = None  # type: ignore

from utils.cli import build_cmd
from ui.components.success_dialog import show_success_with_open

import paramiko

SCADA_IMPORT_ERROR: Optional[str] = None

try:
    from scripts import _Logger as Logger  # type: ignore
    from scripts.functions import sshserver, scada_online  # type: ignore
    from scripts.import_base import company_prefixes  # type: ignore
except Exception as exc_primary:
    try:
        handlers_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), os.pardir))
        project_dir = os.path.abspath(os.path.join(handlers_dir, os.pardir))
        scripts_dir = os.path.join(project_dir, "scripts")
        for candidate in (scripts_dir, project_dir):
            if candidate not in sys.path:
                sys.path.append(candidate)
        import _Logger as Logger  # type: ignore
        from functions import sshserver, scada_online  # type: ignore
        from import_base import company_prefixes  # type: ignore
    except Exception as exc_secondary:
        Logger = None  # type: ignore
        sshserver = None  # type: ignore

        def company_prefixes(_empresa: str) -> Tuple[str, ...]:  # type: ignore
            return tuple()

        info = f"{exc_primary}; {exc_secondary}"
        SCADA_IMPORT_ERROR = info if SCADA_IMPORT_ERROR is None else f"{SCADA_IMPORT_ERROR}; {info}"

try:
    from scripts.hsh_crear_tag import _company_server_candidates, _related_servers  # type: ignore
except Exception:
    def _company_server_candidates(server: str, empresa: str) -> List[str]:  # type: ignore
        return [server]

    def _related_servers(server: str, empresa: str) -> List[str]:  # type: ignore
        return [server]

PI_HOST_MAP = {
    "ITCO": ["itco1his01", "itco1his02"],
    "TRA": ["tra1his01", "tra1his02"],
}

PI_SERVER_MAP = {
    "ITCO": "PI-CO-ITCOTRA01",
    "TRA": "PI-CO-ITCOTRA01",
}

SCADA_HOSTS_FULL = {
    "ITCO": ["itco1sca01", "itco1sca02", "itco1qds01"],
    "TRA": ["tra1sca01", "tra1sca02"],
    "REPS": ["rep1sca01", "rep1sca02", "rep1qds01"],
    "REPP": ["rep2sca01", "rep2sca02", "rep2qds01"],
}

@dataclass
class PiSnapshotResult:
    lines: List[str]
    missing_lines: List[str]
    messages: List[str]
    snapshot_map: Dict[str, List[Dict[str, str]]]
    missing_map: Dict[str, List[str]]
    collected_at: Optional[str]
    has_failures: bool

__all__ = [
    "tk",
    "messagebox",
    "load_workbook",
    "Workbook",
    "build_cmd",
    "show_success_with_open",
    "paramiko",
    "Logger",
    "sshserver",
    "scada_online",
    "company_prefixes",
    "_company_server_candidates",
    "_related_servers",
    "PI_HOST_MAP",
    "PI_SERVER_MAP",
    "SCADA_HOSTS_FULL",
    "SCADA_IMPORT_ERROR",
    "PiSnapshotResult",
    "collect_pi_snapshots",
]

def collect_pi_snapshots(
    *,
    app,
    tags_by_empresa: Dict[str, Set[str]],
    server_map: Dict[str, Optional[str]],
    console_write: Optional[Callable[[str, str], None]] = None,
    record_status: Optional[Callable[[str, str, Optional[bool], Optional[str]], None]] = None,
    empresa_principal: Optional[str] = None,
    derive_prefixes: Optional[Callable[[str], Iterable[str]]] = None,
    pi_server_map: Optional[Dict[str, str]] = None,
    log_filename: str = "pi_query.log",
) -> PiSnapshotResult:
    """
    Ejecuta la consulta de snapshots PI para los tags agrupados por empresa.
    """
    lines: List[str] = []
    missing_lines: List[str] = []
    messages: List[str] = []
    snapshot_map: Dict[str, List[Dict[str, str]]] = {}
    missing_map: Dict[str, List[str]] = {}
    collected_at: Optional[str] = None
    has_failures = False

    console_write = console_write or (lambda msg, tag="info": None)
    empresa_principal_u = empresa_principal.upper() if empresa_principal else None
    pi_server_map = pi_server_map or PI_SERVER_MAP

    def _restore_env(prev: Dict[str, Optional[str]]) -> None:
        for key, value in prev.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value

    def _derive_prefixes_from_hint(hint: Optional[str]) -> List[str]:
        prefixes: List[str] = []
        if hint:
            lowered = hint.lower()
            cut = None
            for marker in ("sca", "qds", "his"):
                idx = lowered.find(marker)
                if idx != -1:
                    cut = idx
                    break
            prefix = hint[:cut] if cut is not None else hint
            if prefix:
                prefixes.append(prefix)
        return prefixes

    def _collect_prefixes(emp_u: str, hint: Optional[str]) -> List[str]:
        prefixes: List[str] = []
        prefixes.extend(_derive_prefixes_from_hint(hint))
        env_sca_hosts = os.environ.get("SCA_HOSTS")
        if env_sca_hosts:
            for host in env_sca_hosts.split(","):
                host = host.strip()
                if host:
                    prefixes.extend(_derive_prefixes_from_hint(host))
        for pref in company_prefixes(emp_u) or []:
            if pref and pref not in prefixes:
                prefixes.append(pref)
        return [pref for pref in prefixes if pref]

    def _encoded_ps_single(emp_u: str, tag: str, pi_server: str, pi_user_ps: str, pi_pass_ps: str) -> str:
        tag_json = json.dumps(tag)
        template = """
$ErrorActionPreference = 'Stop'
Add-Type -Path "C:\\Program Files (x86)\\PIPC\\AF\\PublicAssemblies\\4.0\\OSIsoft.AFSDK.dll"
$piServers = [OSIsoft.AF.PI.PIServers]::GetPIServers()
$piServer = $piServers['{pi_server}']
$securePassword = ConvertTo-SecureString "{pi_pass_ps}" -AsPlainText -Force
$credential = New-Object System.Management.Automation.PSCredential("{pi_user_ps}", $securePassword)
$piServer.Connect($credential)
$target = ConvertFrom-Json @'
{tag_json}
'@
$result = @()
foreach ($piPoint in [OSIsoft.AF.PI.PIPoint]::FindPIPoints($piServer, "{emp_u}*SCADA*", $true)) {{
    if ($piPoint.Name -like "*${{target}}*") {{
        $cv = $piPoint.CurrentValue()
        $result += [PSCustomObject]@{{
            Name = $piPoint.Name
            Value = $cv.Value.ToString()
            Timestamp = $cv.Timestamp.ToString()
        }}
        break
    }}
}}
"__PI_JSON__:" + ($result | ConvertTo-Json -Depth 4 -Compress)
"""
        ps_script = template.format(
            pi_server=pi_server,
            pi_pass_ps=pi_pass_ps,
            pi_user_ps=pi_user_ps,
            tag_json=tag_json,
            emp_u=emp_u,
        )
        encoded = base64.b64encode(ps_script.encode("utf-16-le")).decode("utf-8")
        return f"powershell -NoLogo -NonInteractive -EncodedCommand {encoded}"

    env_pi = app.secure_env() if app else {}
    prev_env = {k: os.environ.get(k) for k in env_pi}
    os.environ.update(env_pi)

    try:
        pi_user = os.environ.get("USER_PI")
        pi_pass = os.environ.get("PASS_PI")
        if not pi_user or not pi_pass:
            msg = "PI: credenciales USER_PI/PASS_PI no disponibles en el entorno."
            lines.append(msg)
            messages.append(msg)
            for emp_candidate in (tags_by_empresa.keys() or server_map.keys()):
                emp_u = str(emp_candidate).upper()
                if record_status:
                    record_status(emp_u, "pi", False, msg)
            return PiSnapshotResult(lines, missing_lines, messages, snapshot_map, missing_map, collected_at, True)

        pi_user_ps = pi_user.replace('"', '`"')
        pi_pass_ps = pi_pass.replace('"', '`"')
        collected_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        log_dir = Path(getattr(app, "base_dir", os.getcwd())) / "out" / "log"
        log_dir.mkdir(parents=True, exist_ok=True)

        try:
            logger, logger_console = Logger.initlog(str(log_dir / log_filename))
        except Exception:
            logger = logger_console = None  # type: ignore

        def _record(emp: str, message: str, ok: Optional[bool]) -> None:
            if record_status:
                record_status(emp, "pi", ok, message)

        for emp_raw, tags in sorted(tags_by_empresa.items(), key=lambda item: item[0]):
            emp_u = str(emp_raw).upper()
            tag_list = sorted({t.strip() for t in tags if t})
            if not tag_list:
                _record(emp_u, "Sin tags generados para consultar", True)
                continue

            if empresa_principal_u and emp_u != empresa_principal_u:
                _record(emp_u, "Valores en PI: no aplica (solo empresa principal)", True)
                continue

            pi_server = pi_server_map.get(emp_u, pi_server_map.get("ITCO", "PI-CO-ITCOTRA01"))
            prefixes = list(derive_prefixes(emp_u)) if derive_prefixes else []
            if not prefixes:
                prefixes = _collect_prefixes(emp_u, server_map.get(emp_u))
            if not prefixes:
                msg = f"PI: no se encontraron prefijos SCADA/HIS para {emp_u}"
                lines.append(msg)
                messages.append(msg)
                has_failures = True
                _record(emp_u, msg, False)
                continue

            aggregated_rows: List[Dict[str, str]] = []
            missing_tags: List[str] = []

            for tag in tag_list:
                tag_rows: Optional[List[Dict[str, str]]] = None
                last_error: Optional[str] = None

                for prefix in prefixes:
                    if tag_rows:
                        break

                    for sca_suffix in ("sca01", "sca02"):
                        sca_host = f"{prefix}{sca_suffix}"
                        console_write(f"[PI] Conectando a SCADA {sca_host} para {tag} ...", "info")
                        client_sca = None
                        try:
                            client_sca = sshserver(sca_host, logger, logger_console)
                        except Exception as exc:
                            last_error = str(exc)
                            console_write(f"[PI] Error conectando a SCADA {sca_host}: {exc}", "error")
                            _record(emp_u, f"{tag}: error conectando a {sca_host}: {exc}", None)
                            continue

                        if client_sca is None:
                            last_error = "sshserver devolvio None"
                            console_write(f"[PI] {sca_host}: sshserver devolvio None", "error")
                            _record(emp_u, f"{tag}: {sca_host} sshserver devolvio None", None)
                            continue

                        transport = client_sca.get_transport()
                        if transport is None or not transport.is_active():
                            last_error = "transporte SSH inactivo"
                            console_write(f"[PI] {sca_host}: transporte SSH inactivo", "error")
                            _record(emp_u, f"{tag}: {sca_host} transporte SSH inactivo", None)
                            try:
                                client_sca.close()
                            except Exception:
                                pass
                            continue

                        try:
                            for his_suffix in ("his01", "his02"):
                                his_host = f"{prefix}{his_suffix}"
                                channel = None
                                client_his = None
                                try:
                                    console_write(f"[PI]   Tunel a {his_host} para {tag} ...", "info")
                                    channel = transport.open_channel("direct-tcpip", (his_host, 22), ("127.0.0.1", 0))
                                    client_his = paramiko.SSHClient()
                                    client_his.set_missing_host_key_policy(paramiko.AutoAddPolicy())
                                    client_his.connect(
                                        hostname=his_host,
                                        username=pi_user,
                                        password=pi_pass,
                                        sock=channel,
                                        look_for_keys=False,
                                    )

                                    command = _encoded_ps_single(emp_u, tag, pi_server, pi_user_ps, pi_pass_ps)
                                    if Logger is not None and logger is not None:
                                        Logger.write_log().log_all('info', f'[PI] comando {command}', logger_console, logger)
                                    stdin, stdout, stderr = client_his.exec_command(command)
                                    output = stdout.read().decode("utf-8", "ignore")
                                    err_output = stderr.read().decode("utf-8", "ignore").strip()
                                    if err_output:
                                        console_write(f"[PI] {his_host} STDERR: {err_output}", "warn")

                                    marker = "__PI_JSON__:"
                                    idx = output.find(marker)
                                    if idx == -1:
                                        last_error = err_output or output.strip() or "respuesta inesperada"
                                        console_write(f"[PI] {his_host}: respuesta inesperada -> {last_error}", "warn")
                                        _record(emp_u, f"{tag}: {his_host} respuesta inesperada", None)
                                        continue

                                    json_text = output[idx + len(marker):].strip()
                                    raw_rows = json.loads(json_text) if json_text else []

                                    normalized: List[Dict[str, str]] = []
                                    if isinstance(raw_rows, list):
                                        for item in raw_rows:
                                            if isinstance(item, dict):
                                                normalized.append({
                                                    "Name": str(item.get("Name", item.get("name", "?"))),
                                                    "Value": str(item.get("Value", item.get("value", ""))),
                                                    "Timestamp": str(item.get("Timestamp", item.get("timestamp", ""))),
                                                })
                                    elif isinstance(raw_rows, dict):
                                        normalized.append({
                                            "Name": str(raw_rows.get("Name", raw_rows.get("name", "?"))),
                                            "Value": str(raw_rows.get("Value", raw_rows.get("value", ""))),
                                            "Timestamp": str(raw_rows.get("Timestamp", raw_rows.get("timestamp", ""))),
                                        })

                                    if normalized:
                                        tag_rows = normalized
                                        aggregated_rows.extend(normalized)
                                        _record(emp_u, f"{tag}: {len(normalized)} valores desde {his_host}", None)
                                        break

                                    last_error = "sin valores"
                                    console_write(f"[PI] {his_host}: sin valores para {tag}", "warn")
                                except Exception as exc:
                                    last_error = str(exc)
                                    console_write(f"[PI] Error {sca_host}->{his_host} ({tag}): {exc}", "error")
                                    _record(emp_u, f"{tag}: error {sca_host}->{his_host}: {exc}", None)
                                finally:
                                    if client_his is not None:
                                        client_his.close()
                                    if channel is not None:
                                        channel.close()
                                if tag_rows:
                                    break
                        finally:
                            try:
                                if client_sca is not None:
                                    client_sca.close()
                            except Exception:
                                pass

                        if tag_rows:
                            break

                if not tag_rows:
                    missing_tags.append(tag)
                    has_failures = True
                    warning = f"{emp_u}:{tag}:{last_error or 'sin datos'}"
                    missing_lines.append(warning)
                    console_write(f"[PI] {emp_u}: no se obtuvo informacion para {tag} ({last_error or 'sin datos'})", "warn")
                    _record(emp_u, f"{tag}: {last_error or 'sin datos'}", False)
                else:
                    console_write(f"[PI] {emp_u}: valores capturados para {tag}", "info")

            snapshot_map[emp_u] = aggregated_rows
            missing_map[emp_u] = missing_tags

            if aggregated_rows:
                pretty = ", ".join(f"{row.get('Name', '?')} = {row.get('Value', '?')}" for row in aggregated_rows)
                msg = f"{emp_u}: {pretty}"
                lines.append(msg)
                messages.append(msg)
                if record_status and not missing_tags:
                    record_status(emp_u, "pi", True, msg)
            if missing_tags:
                miss_msg = f"{emp_u}: sin datos para {', '.join(missing_tags)}"
                lines.append(miss_msg)
                messages.append(miss_msg)

            if not aggregated_rows and not missing_tags:
                _record(emp_u, "Sin valores consultados", True)

        return PiSnapshotResult(lines, missing_lines, messages, snapshot_map, missing_map, collected_at, has_failures)
    finally:
        _restore_env(prev_env)
